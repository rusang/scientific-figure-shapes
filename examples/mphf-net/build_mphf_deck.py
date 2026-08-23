#!/usr/bin/env python3
"""Rebuild assets/reference-figures/mphf-net-architecture.png as an editable PPTX.

源图 1536x1024 px -> slide 960x640 pt（uniform 0.625 pt/px）。
可编辑元素全部重建；PCB 板图按 Crop Contract 保留为 assets/01_R1_pcb.png。
"""
from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_LINE_DASH_STYLE as MSO_LINE
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt

HERE = Path(__file__).parent
SCALE = 0.625  # pt per source px

NAVY = RGBColor(24, 33, 74)
BLUE = RGBColor(37, 80, 216)
RED = RGBColor(224, 58, 38)
PURPLE = RGBColor(112, 62, 200)
GREEN = RGBColor(28, 148, 84)
INK = RGBColor(40, 44, 60)
GRAY = RGBColor(120, 126, 140)
PANEL_EDGE = RGBColor(214, 219, 232)
PANEL_BG = RGBColor(252, 252, 254)
CUBE_BLUE = (108, 146, 232)
CUBE_GREEN = (124, 198, 132)
CUBE_ORANGE = (240, 126, 70)
BADGE_UP = RGBColor(196, 172, 240)
BADGE_DOWN = RGBColor(244, 172, 96)
BADGE_PLUS = RGBColor(52, 58, 78)
GRID_BG = RGBColor(214, 228, 252)
GRID_LINE = RGBColor(120, 156, 228)


def px(value: float) -> Pt:
    return Pt(value * SCALE)


def rect(slide, name, x, y, w, h, *, fill=None, edge=None, edge_w=1.0,
         shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.08, dash=None):
    obj = slide.shapes.add_shape(shape, px(x), px(y), px(w), px(h))
    obj.name = name
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        try:
            obj.adjustments[0] = radius
        except (IndexError, ValueError):
            pass
    if fill is None:
        obj.fill.background()
    else:
        obj.fill.solid()
        obj.fill.fore_color.rgb = fill
    if edge is None:
        obj.line.fill.background()
    else:
        obj.line.color.rgb = edge
        obj.line.width = Pt(edge_w)
        if dash:
            obj.line.dash_style = dash
    obj.shadow.inherit = False
    obj.text_frame.word_wrap = True
    return obj


def label(slide, name, x, y, w, h, text, *, size=8, color=INK, bold=False,
          align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE):
    box = slide.shapes.add_textbox(px(x), px(y), px(w), px(h))
    box.name = name
    frame = box.text_frame
    frame.word_wrap = True
    frame.vertical_anchor = anchor
    frame.margin_left = frame.margin_right = Emu(9525)
    frame.margin_top = frame.margin_bottom = Emu(4762)
    first = True
    for line in text.split("\n"):
        paragraph = frame.paragraphs[0] if first else frame.add_paragraph()
        first = False
        paragraph.alignment = align
        run = paragraph.add_run()
        run.text = line
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = color
    return box


def _line_style(shape, color, width, dash=False, arrow=True):
    shape.line.color.rgb = color
    shape.line.width = Pt(width)
    element = shape.line._get_or_add_ln()
    if dash:
        pd = element.makeelement(qn("a:prstDash"), {"val": "dash"})
        element.append(pd)
    if arrow:
        tail = element.makeelement(
            qn("a:tailEnd"), {"type": "triangle", "w": "med", "len": "med"}
        )
        element.append(tail)


def line(slide, name, x1, y1, x2, y2, *, color=INK, width=1.4, dash=False,
         arrow=True):
    shape = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT, px(x1), px(y1), px(x2), px(y2)
    )
    shape.name = name
    _line_style(shape, color, width, dash=dash, arrow=arrow)
    return shape


def badge(slide, name, cx, cy, r, fill, text, *, size=6.5,
          text_color=RGBColor(255, 255, 255)):
    obj = slide.shapes.add_shape(MSO_SHAPE.OVAL, px(cx - r), px(cy - r),
                                 px(2 * r), px(2 * r))
    obj.name = name
    obj.fill.solid()
    obj.fill.fore_color.rgb = fill
    obj.line.fill.background()
    obj.shadow.inherit = False
    label(slide, f"{name}_text", cx - r, cy - r, 2 * r, 2 * r, text,
          size=size, color=text_color, bold=True)
    return obj


def cuboid(slide, prefix, x, y, w, h, base, depth=8):
    """三面 cuboid：front/top/right，top 亮 side 暗，统一右上深度向量。"""
    r, g, b = base
    front = RGBColor(r, g, b)
    top = RGBColor(min(r + 28, 255), min(g + 28, 255), min(b + 28, 255))
    right = RGBColor(max(r - 44, 0), max(g - 44, 0), max(b - 44, 0))
    rect(slide, f"{prefix}_top", x, y, w, depth, fill=top,
         shape=MSO_SHAPE.RECTANGLE)
    rect(slide, f"{prefix}_front", x, y + depth, w, h, fill=front,
         shape=MSO_SHAPE.RECTANGLE)
    rect(slide, f"{prefix}_right", x + w, y + depth, depth, h, fill=right,
         shape=MSO_SHAPE.RECTANGLE)


def grid(slide, prefix, x, y, w, h, cols=4, rows=3):
    rect(slide, f"{prefix}_bg", x, y, w, h, fill=GRID_BG, edge=GRID_LINE,
         edge_w=0.75, shape=MSO_SHAPE.RECTANGLE)
    for col in range(1, cols):
        line(slide, f"{prefix}_v{col}", x + w * col / cols, y,
             x + w * col / cols, y + h, color=GRID_LINE, width=0.5,
             arrow=False)
    for row in range(1, rows):
        line(slide, f"{prefix}_h{row}", x, y + h * row / rows, x + w,
             y + h * row / rows, color=GRID_LINE, width=0.5, arrow=False)


def square_row(slide, prefix, x, y, count, size=13, gap=6,
               fill=RGBColor(150, 120, 220)):
    for index in range(count):
        rect(slide, f"{prefix}_{index}", x + index * (size + gap), y, size,
             size, fill=fill, shape=MSO_SHAPE.RECTANGLE, radius=0)
    label(slide, f"{prefix}_dots", x + count * (size + gap) + 2, y - 2, 30,
          size + 4, "····", size=8, color=GRAY, bold=True)


def mini_bars(slide, prefix, x, y):
    heights = (18, 30, 12, 24)
    for index, height in enumerate(heights):
        rect(slide, f"{prefix}_{index}", x + index * 10, y + 32 - height, 7,
             height, fill=RGBColor(90, 118, 216), shape=MSO_SHAPE.RECTANGLE,
             radius=0)


def build() -> Path:
    presentation = Presentation()
    presentation.slide_width = px(1536)
    presentation.slide_height = px(1024)
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])

    # ---- 标题横幅 ----
    rect(slide, "SUMMER_B_banner", 10, 8, 1516, 54, fill=NAVY, radius=0.18)
    label(slide, "SUMMER_T_banner", 10, 12, 1516, 46,
          "MPHF-Net 工业表面缺陷检测网络结构设计图", size=17,
          color=RGBColor(255, 255, 255), bold=True)

    # ---- 顶行面板底框 ----
    panels = {
        "input": (12, 76, 200, 460),
        "backbone": (228, 76, 258, 460),
        "neck": (498, 76, 574, 460),
        "head": (1086, 76, 240, 460),
        "legend": (1340, 76, 186, 460),
    }
    for key, (x, y, w, h) in panels.items():
        rect(slide, f"SUMMER_E_panel_{key}", x, y, w, h, fill=PANEL_BG,
             edge=PANEL_EDGE, edge_w=1.2, radius=0.05)

    # ---- P1 输入图像 ----
    label(slide, "SUMMER_T_input_title", 12, 86, 200, 22, "输入图像",
          size=11, bold=True)
    pic = slide.shapes.add_picture(str(HERE / "assets/01_R1_pcb.png"),
                                   px(25), px(162), px(171), px(244))
    pic.name = "SUMMER_R1_pcb"
    label(slide, "SUMMER_T_input_size", 12, 430, 200, 24, "640 × 640 × 3",
          size=11, color=INK)

    # ---- P2 Backbone ----
    label(slide, "SUMMER_T_backbone_title", 228, 86, 258, 20,
          "Backbone：CSP-MEEM", size=11, color=BLUE, bold=True)
    label(slide, "SUMMER_T_backbone_sub", 228, 108, 258, 16,
          "多尺度细粒度特征提取", size=8, color=GRAY)
    pyramid = [  # (x, y, w, h) 自上而下渐大，相邻层叠（front 与下一层 top 相交）
        (290, 150, 36, 30),
        (270, 182, 52, 46),
        (248, 232, 70, 64),
        (222, 296, 96, 84),
    ]
    for index, (x, y, w, h) in enumerate(pyramid, start=1):
        cuboid(slide, f"SUMMER_E_bb{index}", x, y, w, h, CUBE_BLUE)
    label(slide, "SUMMER_T_backbone_dots", 270, 452, 40, 18, "⋮", size=11,
          color=GRAY, bold=True)
    # 箭头从对应方块（bb2/bb3/bb4）右缘中心水平指向各 stage 框左缘
    stage = [("C5", "20×C", BLUE, pyramid[1]), ("C4", "40×C", BLUE,
             pyramid[2]), ("C3", "80×C", RED, pyramid[3])]
    for name, dim, color, (bx, by, bw, bh) in stage:
        center_y = by + 8 + bh / 2
        rect(slide, f"SUMMER_E_bb_{name.lower()}", 396, center_y - 22, 58,
             44, fill=RGBColor(245, 248, 255), edge=color, edge_w=1.4)
        label(slide, f"SUMMER_T_bb_{name.lower()}", 396, center_y - 18, 58,
              20, name, size=10, color=color, bold=True)
        label(slide, f"SUMMER_T_bb_{name.lower()}_dim", 396, center_y, 58,
              18, dim, size=8, color=color)
        line(slide, f"SUMMER_L_bb_to_{name.lower()}", bx + bw + 8, center_y,
             396, center_y, color=INK, width=1.2)
    rect(slide, "SUMMER_E_bb_outbar", 240, 486, 216, 30,
         fill=RGBColor(235, 241, 255), edge=PANEL_EDGE)
    label(slide, "SUMMER_T_bb_outbar", 240, 490, 216, 22, "输出多尺度特征",
          size=9, color=INK)

    # ---- P3 Neck ----
    label(slide, "SUMMER_T_neck_title", 498, 86, 574, 20,
          "Neck：FPN-PAN （LOSC + CGAFusion）", size=11, color=GREEN,
          bold=True)
    rect(slide, "SUMMER_E_neck_td", 540, 116, 130, 26,
         fill=RGBColor(226, 244, 230), edge=GREEN, edge_w=0.9)
    label(slide, "SUMMER_T_neck_td", 540, 119, 130, 20, "Top-Down 路径",
          size=8.5, color=GREEN, bold=True)
    rect(slide, "SUMMER_E_neck_bu", 828, 116, 130, 26,
         fill=RGBColor(252, 236, 226), edge=RGBColor(226, 120, 60),
         edge_w=0.9)
    label(slide, "SUMMER_T_neck_bu", 828, 119, 130, 20, "Bottom-Up 路径",
          size=8.5, color=RGBColor(206, 100, 40), bold=True)

    def neck_tag(name, x, y, text, color):
        rect(slide, f"SUMMER_E_{name}", x, y, 40, 28,
             fill=RGBColor(245, 248, 255), edge=color, edge_w=1.2)
        label(slide, f"SUMMER_T_{name}", x, y + 4, 40, 20, text, size=9,
              color=color, bold=True)

    neck_tag("neck_c5", 552, 152, "C5", BLUE)
    neck_tag("neck_c4", 508, 266, "C4", BLUE)
    neck_tag("neck_c3", 508, 408, "C3", BLUE)

    badge(slide, "SUMMER_E_neck_up1", 602, 222, 14, BADGE_UP, "UP",
          text_color=PURPLE)
    badge(slide, "SUMMER_E_neck_up2", 602, 344, 14, BADGE_UP, "UP",
          text_color=PURPLE)
    badge(slide, "SUMMER_E_neck_add_td1", 602, 280, 13, BADGE_PLUS, "+")
    badge(slide, "SUMMER_E_neck_add_td2", 602, 420, 13, BADGE_PLUS, "+")

    def fusion(name, x, y):
        rect(slide, f"SUMMER_E_{name}", x, y, 92, 30,
             fill=RGBColor(240, 245, 255), edge=BLUE, edge_w=1.2)
        label(slide, f"SUMMER_T_{name}", x, y + 5, 92, 20, "CGAFusion",
              size=8.5, color=BLUE, bold=True)

    fusion("neck_fusion1", 660, 265)
    fusion("neck_fusion2", 660, 405)

    badge(slide, "SUMMER_E_neck_add_p5", 892, 182, 13, BADGE_PLUS, "+")
    badge(slide, "SUMMER_E_neck_add_p4", 892, 342, 13, BADGE_PLUS, "+")
    badge(slide, "SUMMER_E_neck_add_p3", 892, 468, 13, BADGE_PLUS, "+")
    badge(slide, "SUMMER_E_neck_down1", 892, 242, 14, BADGE_DOWN, "Down",
          size=5, text_color=RGBColor(120, 60, 10))
    badge(slide, "SUMMER_E_neck_down2", 892, 404, 14, BADGE_DOWN, "Down",
          size=5, text_color=RGBColor(120, 60, 10))
    rect(slide, "SUMMER_E_neck_losc", 856, 288, 72, 26,
         fill=RGBColor(253, 240, 232), edge=RGBColor(226, 110, 50),
         edge_w=1.2)
    label(slide, "SUMMER_T_neck_losc", 856, 291, 72, 20, "LOSC", size=9,
          color=RGBColor(206, 90, 30), bold=True)

    pyr = [("p5", 964, 158, (20, "P5", "20×C", CUBE_BLUE)),
           ("p4", 964, 318, (20, "P4", "40×C", CUBE_GREEN)),
           ("p3", 964, 444, (20, "P3", "80×C", CUBE_ORANGE))]
    for key, x, y, (_, name, dim, base) in pyr:
        cuboid(slide, f"SUMMER_E_neck_{key}", x, y, 34, 26, base)
        color = {"P5": BLUE, "P4": GREEN, "P3": RED}[name]
        label(slide, f"SUMMER_T_neck_{key}", x + 48, y - 4, 52, 20, name,
              size=10, color=color, bold=True)
        label(slide, f"SUMMER_T_neck_{key}_dim", x + 48, y + 14, 52, 16, dim,
              size=8, color=color)

    # Neck 主干箭头（直线，审计目标）
    # 箭头端点贴目标边界（badge 字标已列入 routing manifest 的
    # ignore_text_shapes 窄豁免，线触圆边不再退让悬空）
    line(slide, "SUMMER_L_neck_c5_up", 572, 182, 601, 208)
    line(slide, "SUMMER_L_neck_up1_add", 602, 236, 602, 267)
    line(slide, "SUMMER_L_neck_c4_add", 548, 280, 589, 280)
    line(slide, "SUMMER_L_neck_add1_fusion", 615, 280, 660, 280)
    line(slide, "SUMMER_L_neck_up2_add", 602, 358, 602, 407)
    line(slide, "SUMMER_L_neck_c3_add", 548, 421, 589, 421)
    line(slide, "SUMMER_L_neck_add2_fusion", 615, 420, 660, 420)
    line(slide, "SUMMER_L_neck_fusion1_add", 752, 280, 881, 189)
    line(slide, "SUMMER_L_neck_fusion2_add", 752, 420, 881, 461)
    line(slide, "SUMMER_L_neck_addp5_cube", 905, 182, 964, 181)
    line(slide, "SUMMER_L_neck_addp4_cube", 905, 342, 964, 341)
    line(slide, "SUMMER_L_neck_addp3_cube", 905, 468, 964, 467)
    line(slide, "SUMMER_L_neck_addp5_down", 892, 195, 892, 228)
    line(slide, "SUMMER_L_neck_down_losc", 892, 256, 892, 288)
    line(slide, "SUMMER_L_neck_losc_addp4", 892, 314, 892, 329)
    line(slide, "SUMMER_L_neck_addp4_down2", 892, 355, 892, 390)
    line(slide, "SUMMER_L_neck_down2_addp3", 892, 418, 892, 455)
    # 跨层虚线（装饰，不入审计前缀）
    line(slide, "SUMMER_LD_skip1", 706, 296, 706, 405, dash=True,
         color=GRAY, width=1.0)
    line(slide, "SUMMER_LD_skip2", 620, 152, 660, 272, dash=True,
         color=GRAY, width=1.0, arrow=False)

    rect(slide, "SUMMER_E_neck_note_cga", 506, 508, 250, 28,
         fill=RGBColor(235, 242, 255), edge=BLUE, edge_w=0.8)
    label(slide, "SUMMER_T_neck_note_cga", 506, 512, 250, 20,
          "CGAFusion（高低频自适应融合）", size=8.5, color=BLUE)
    rect(slide, "SUMMER_E_neck_note_losc", 774, 508, 250, 28,
         fill=RGBColor(253, 238, 232), edge=RED, edge_w=0.8)
    label(slide, "SUMMER_T_neck_note_losc", 774, 512, 250, 20,
          "LOSC（方向性下采样建模）", size=8.5, color=RED)

    # ---- P4 Head ----
    label(slide, "SUMMER_T_head_title", 1086, 86, 240, 20,
          "Head：RT-DETR 原始检测头", size=10.5, color=PURPLE, bold=True)
    rect(slide, "SUMMER_E_head_iou", 1098, 142, 216, 96,
         fill=RGBColor(255, 255, 255), edge=PANEL_EDGE, edge_w=1.0)
    label(slide, "SUMMER_T_head_iou", 1098, 152, 216, 36,
          "IoU-Aware\nQuery Selection", size=9, color=INK, bold=True)
    square_row(slide, "SUMMER_E_head_iou_sq", 1116, 202, 7, size=12, gap=5)
    line(slide, "SUMMER_L_head_iou_dec", 1206, 238, 1206, 292)
    rect(slide, "SUMMER_E_head_dec", 1098, 292, 216, 108,
         fill=RGBColor(255, 255, 255), edge=PANEL_EDGE, edge_w=1.0)
    label(slide, "SUMMER_T_head_dec", 1098, 300, 216, 22,
          "Transformer Decoder × 6 层", size=9, color=INK, bold=True)
    square_row(slide, "SUMMER_E_head_dec_sq", 1114, 340, 6, size=14, gap=7)
    line(slide, "SUMMER_L_head_dec_out", 1206, 400, 1206, 432)
    rect(slide, "SUMMER_E_head_out", 1098, 432, 216, 92,
         fill=RGBColor(255, 255, 255), edge=PANEL_EDGE, edge_w=1.0)
    label(slide, "SUMMER_T_head_out", 1098, 438, 216, 34,
          "输出：\n类别概率 + 归一化边界框坐标", size=8.5, color=INK)
    mini_bars(slide, "SUMMER_E_head_bars", 1136, 478)
    rect(slide, "SUMMER_E_head_bbox", 1216, 478, 40, 32, fill=None,
         edge=BLUE, edge_w=1.2, shape=MSO_SHAPE.RECTANGLE,
         dash=MSO_LINE.DASH)

    # ---- P5 图例 ----
    label(slide, "SUMMER_T_legend_title", 1340, 86, 186, 20, "图例说明",
          size=10.5, bold=True)
    cubes = [("lg_p5", "P5 / C5 （20×C）", CUBE_BLUE, 132),
             ("lg_p4", "P4 / C4 （40×C）", CUBE_GREEN, 168),
             ("lg_p3", "P3 / C3 （80×C）", CUBE_ORANGE, 204)]
    for key, text, base, y in cubes:
        cuboid(slide, f"SUMMER_E_{key}", 1354, y, 20, 14, base, depth=5)
        label(slide, f"SUMMER_T_{key}", 1388, y - 2, 136, 22, text, size=7.5,
              align=PP_ALIGN.LEFT)
    badge(slide, "SUMMER_E_lg_up", 1364, 250, 10, BADGE_UP, "UP",
          text_color=PURPLE, size=5.5)
    label(slide, "SUMMER_T_lg_up", 1388, 238, 136, 22, "上采样 （× 2）",
          size=7.5, align=PP_ALIGN.LEFT)
    badge(slide, "SUMMER_E_lg_down", 1364, 286, 10, BADGE_DOWN, "Down",
          text_color=RGBColor(120, 60, 10), size=4.5)
    label(slide, "SUMMER_T_lg_down", 1388, 274, 136, 22, "下采样 （× 2）",
          size=7.5, align=PP_ALIGN.LEFT)
    badge(slide, "SUMMER_E_lg_add", 1364, 322, 10, BADGE_PLUS, "+", size=7)
    label(slide, "SUMMER_T_lg_add", 1388, 310, 136, 22, "拼接 （Concat）",
          size=7.5, align=PP_ALIGN.LEFT)
    rect(slide, "SUMMER_E_lg_losc", 1352, 350, 32, 16,
         fill=RGBColor(253, 240, 232), edge=RGBColor(226, 110, 50),
         edge_w=0.8)
    label(slide, "SUMMER_T_lg_losc_tag", 1352, 351, 32, 14, "LOSC", size=5.5,
          color=RGBColor(206, 90, 30), bold=True)
    label(slide, "SUMMER_T_lg_losc", 1388, 347, 136, 22, "方向性下采样建模",
          size=7.5, align=PP_ALIGN.LEFT)
    rect(slide, "SUMMER_E_lg_cga", 1352, 384, 32, 16,
         fill=RGBColor(235, 242, 255), edge=BLUE, edge_w=0.8)
    label(slide, "SUMMER_T_lg_cga_tag", 1352, 385, 32, 14, "CGAFusion",
          size=4, color=BLUE, bold=True)
    label(slide, "SUMMER_T_lg_cga", 1388, 381, 136, 22, "高低频自适应融合",
          size=7.5, align=PP_ALIGN.LEFT)
    line(slide, "SUMMER_LD_lg_solid", 1354, 428, 1384, 428)
    label(slide, "SUMMER_T_lg_solid", 1388, 417, 136, 22, "数据流向",
          size=7.5, align=PP_ALIGN.LEFT)
    line(slide, "SUMMER_LD_lg_dash", 1354, 462, 1384, 462, dash=True)
    label(slide, "SUMMER_T_lg_dash", 1388, 451, 136, 22, "跨层连接",
          size=7.5, align=PP_ALIGN.LEFT)

    # ---- 顶行主流程箭头 ----
    line(slide, "SUMMER_L_input_backbone", 212, 296, 228, 296, color=BLUE,
         width=4.0)
    line(slide, "SUMMER_L_backbone_neck", 486, 296, 498, 296, color=BLUE,
         width=4.0)
    line(slide, "SUMMER_L_neck_head", 1072, 296, 1086, 296, color=BLUE,
         width=4.0)

    # ---- 中行三模块 ----
    modules = {
        "m1": (12, 556, 474, 284),
        "m2": (498, 556, 408, 284),
        "m3": (916, 556, 610, 284),
    }
    for key, (x, y, w, h) in modules.items():
        rect(slide, f"SUMMER_E_panel_{key}", x, y, w, h, fill=PANEL_BG,
             edge=PANEL_EDGE, edge_w=1.2, radius=0.04)
    label(slide, "SUMMER_T_m1_title", 12, 566, 474, 20,
          "模块一：CSP-MEEM（多尺度细粒度特征提取）", size=10, color=BLUE,
          bold=True)
    flow = [("m1_in", "输入图像", 600), ("m1_csp", "CSP\n分割", 636),
            ("m1_branch", "分支1  分支2  ···  分支n", 690),
            ("m1_fuse", "融合", 736)]
    label(slide, "SUMMER_T_m1_in", 60, 596, 120, 20, "输入图像", size=8)
    rect(slide, "SUMMER_E_m1_csp", 84, 622, 72, 30,
         fill=RGBColor(235, 242, 255), edge=BLUE, edge_w=0.9)
    label(slide, "SUMMER_T_m1_csp", 84, 626, 72, 22, "CSP 分割", size=8,
          color=INK)
    for index, name in enumerate(("分支1", "分支2", "分支n")):
        x = 28 + index * 66
        rect(slide, f"SUMMER_E_m1_b{index}", x, 674, 52, 26,
             fill=RGBColor(240, 246, 255), edge=BLUE, edge_w=0.8)
        label(slide, f"SUMMER_T_m1_b{index}", x, 677, 52, 20, name, size=7.5)
    label(slide, "SUMMER_T_m1_bdots", 156, 677, 30, 20, "···", size=8,
          color=GRAY)
    rect(slide, "SUMMER_E_m1_fuse", 70, 716, 100, 26,
         fill=RGBColor(235, 242, 255), edge=BLUE, edge_w=0.9)
    label(slide, "SUMMER_T_m1_fuse", 70, 719, 100, 20, "融合", size=8)
    rect(slide, "SUMMER_E_m1_ema", 62, 754, 116, 26,
         fill=RGBColor(255, 240, 226), edge=RGBColor(226, 130, 50),
         edge_w=0.9)
    label(slide, "SUMMER_T_m1_ema", 62, 757, 116, 20, "EMA 注意力", size=8,
          color=RGBColor(200, 100, 30))
    label(slide, "SUMMER_T_m1_out", 60, 790, 120, 20, "输出特征", size=8)
    line(slide, "SUMMER_L_m1_f1", 120, 616, 120, 622, width=1.0)
    line(slide, "SUMMER_L_m1_f2", 120, 652, 120, 674, width=1.0)
    line(slide, "SUMMER_L_m1_f3", 120, 700, 120, 716, width=1.0)
    line(slide, "SUMMER_L_m1_f4", 120, 742, 120, 754, width=1.0)
    line(slide, "SUMMER_L_m1_f5", 120, 780, 120, 790, width=1.0)
    m1_bullets = ("多尺度并行感知，捕获不同\n大小的缺陷信息",
                  "增强细节特征表达，保留\n边缘和纹理",
                  "EMA 注意力增强关键特征，\n抑制冗余信息")
    for index, text in enumerate(m1_bullets):
        label(slide, f"SUMMER_T_m1_p{index}", 226, 606 + index * 66, 246, 56,
              "•  " + text, size=8.5, align=PP_ALIGN.LEFT,
              anchor=MSO_ANCHOR.TOP)

    label(slide, "SUMMER_T_m2_title", 498, 566, 408, 20,
          "模块二：LOSC（方向性下采样建模）", size=10, color=RED, bold=True)
    label(slide, "SUMMER_T_m2_in", 520, 594, 120, 18, "输入特征", size=8)
    grid(slide, "SUMMER_G_m2_in", 540, 614, 76, 56)
    badge(slide, "SUMMER_E_m2_down", 578, 706, 15, BADGE_DOWN, "Down",
          size=5.5, text_color=RGBColor(120, 60, 10))
    grid(slide, "SUMMER_G_m2_out", 552, 748, 52, 40, cols=3, rows=2)
    label(slide, "SUMMER_T_m2_out", 520, 794, 120, 18, "输出特征", size=8)
    line(slide, "SUMMER_L_m2_f1", 578, 670, 578, 690, width=1.0)
    line(slide, "SUMMER_L_m2_f2", 578, 722, 578, 748, width=1.0)
    m2_bullets = ("引入方向性卷积核", "沿关键方向提取特征",
                  "抑制无关干扰，保留方向信息",
                  "增强对工业表面纹理与缺陷\n形态的建模能力")
    for index, text in enumerate(m2_bullets):
        label(slide, f"SUMMER_T_m2_p{index}", 660, 600 + index * 52, 236, 48,
              "•  " + text, size=8.5, align=PP_ALIGN.LEFT,
              anchor=MSO_ANCHOR.TOP)

    label(slide, "SUMMER_T_m3_title", 916, 566, 610, 20,
          "模块三：CGAFusion（高低频自适应融合）", size=10, color=BLUE,
          bold=True)
    label(slide, "SUMMER_T_m3_fh", 930, 598, 130, 18, "输入特征 Fh（高频）",
          size=7.5, align=PP_ALIGN.LEFT)
    grid(slide, "SUMMER_G_m3_fh", 946, 618, 64, 46)
    rect(slide, "SUMMER_E_m3_high", 1046, 620, 96, 40,
         fill=RGBColor(235, 242, 255), edge=BLUE, edge_w=0.9)
    label(slide, "SUMMER_T_m3_high", 1046, 624, 96, 32, "高频分支\n（细节信息）",
          size=7.5, color=BLUE)
    label(slide, "SUMMER_T_m3_fl", 930, 682, 130, 18, "输入特征 Fl（低频）",
          size=7.5, align=PP_ALIGN.LEFT)
    grid(slide, "SUMMER_G_m3_fl", 946, 702, 64, 46)
    rect(slide, "SUMMER_E_m3_low", 1046, 704, 96, 40,
         fill=RGBColor(235, 248, 238), edge=GREEN, edge_w=0.9)
    label(slide, "SUMMER_T_m3_low", 1046, 708, 96, 32, "低频分支\n（结构信息）",
          size=7.5, color=GREEN)
    badge(slide, "SUMMER_E_m3_add", 1190, 682, 14, BADGE_PLUS, "+")
    grid(slide, "SUMMER_G_m3_out", 1166, 730, 52, 40, cols=3, rows=2)
    label(slide, "SUMMER_T_m3_out", 1140, 776, 110, 18, "输出特征", size=8)
    line(slide, "SUMMER_L_m3_f1", 1142, 640, 1176, 672, width=1.0)
    line(slide, "SUMMER_L_m3_f2", 1142, 724, 1176, 692, width=1.0)
    line(slide, "SUMMER_L_m3_f3", 1190, 696, 1190, 730, width=1.0)
    m3_bullets = ("分离高频细节与低频结构信息", "空间联合注意力自适应融合",
                  "突出关键细节，抑制噪声干扰", "提升特征表达与泛化能力")
    for index, text in enumerate(m3_bullets):
        label(slide, f"SUMMER_T_m3_p{index}", 1268, 600 + index * 52, 250,
              44, "•  " + text, size=8.5, align=PP_ALIGN.LEFT,
              anchor=MSO_ANCHOR.TOP)

    # ---- 底行四面板 ----
    bottoms = {
        "b1": (12, 852, 322, 164),
        "b2": (348, 852, 542, 164),
        "b3": (908, 852, 226, 164),
        "b4": (1148, 852, 378, 164),
    }
    for key, (x, y, w, h) in bottoms.items():
        rect(slide, f"SUMMER_E_panel_{key}", x, y, w, h, fill=PANEL_BG,
             edge=PANEL_EDGE, edge_w=1.2, radius=0.06)
    label(slide, "SUMMER_T_b1_title", 12, 860, 322, 20,
          "输出多尺度特征金字塔", size=10, color=BLUE, bold=True)
    bottom_pyr = [("SUMMER_E_out_p3", 46, 896, 62, 48, CUBE_ORANGE,
                   "P3\n（80×C）", RED),
                  ("SUMMER_E_out_p4", 142, 904, 50, 40, CUBE_GREEN,
                   "P4\n（40×C）", GREEN),
                  ("SUMMER_E_out_p5", 226, 912, 40, 32, CUBE_BLUE,
                   "P5\n（20×C）", BLUE)]
    for name, x, y, w, h, base, text, color in bottom_pyr:
        cuboid(slide, name, x, y, w, h, base)
        label(slide, f"{name}_label", x - 12, y + h + 14, w + 34, 32, text,
              size=8, color=color, bold=True)
    label(slide, "SUMMER_T_b2_title", 348, 860, 542, 20,
          "输入到 RT-DETR Head", size=10, color=INK, bold=True)
    label(slide, "SUMMER_T_b2_iou", 380, 892, 130, 34,
          "IoU-Aware\nQuery Selection", size=8, color=INK)
    square_row(slide, "SUMMER_E_b2_sq1", 386, 934, 6, size=12, gap=5)
    label(slide, "SUMMER_T_b2_dec", 560, 880, 240, 18,
          "Transformer Decoder × 6 层", size=8.5, color=INK)
    square_row(slide, "SUMMER_E_b2_sq2", 584, 922, 5, size=16, gap=9,
               fill=RGBColor(140, 110, 214))
    label(slide, "SUMMER_T_b2_dots2", 742, 922, 40, 20, "····", size=9,
          color=GRAY, bold=True)
    rect(slide, "SUMMER_E_b2_last", 792, 922, 16, 16,
         fill=RGBColor(140, 110, 214), shape=MSO_SHAPE.RECTANGLE, radius=0)
    line(slide, "SUMMER_L_b2_flow", 516, 930, 584, 930, width=1.2)
    line(slide, "SUMMER_L_b1_b2", 334, 934, 348, 934, color=BLUE, width=2.5)
    line(slide, "SUMMER_L_b2_b3", 890, 934, 908, 934, color=BLUE, width=2.5)
    label(slide, "SUMMER_T_b3_title", 908, 860, 226, 20, "检测结果", size=10,
          color=INK, bold=True)
    mini_bars(slide, "SUMMER_E_b3_bars", 936, 896)
    label(slide, "SUMMER_T_b3_bars", 916, 946, 100, 18, "类别概率", size=7.5)
    rect(slide, "SUMMER_E_b3_box1", 1030, 892, 46, 36, fill=None, edge=RED,
         edge_w=1.2, shape=MSO_SHAPE.RECTANGLE)
    rect(slide, "SUMMER_E_b3_box2", 1048, 908, 46, 36, fill=None,
         edge=GREEN, edge_w=1.2, shape=MSO_SHAPE.RECTANGLE)
    label(slide, "SUMMER_T_b3_box", 1010, 946, 120, 30, "归一化边界\n框坐标",
          size=7.5)
    label(slide, "SUMMER_T_b4_title", 1148, 860, 378, 20, "三模块协同优势",
          size=10, color=BLUE, bold=True)
    advantages = [("更精细：", "保留小目标细粒度特征", BLUE),
                  ("更鲁棒：", "增强方向性缺陷感知", GREEN),
                  ("更高效：", "自适应融合多尺度特征", RED)]
    for index, (head, tail, color) in enumerate(advantages):
        y = 890 + index * 40
        badge(slide, f"SUMMER_E_b4_icon{index}", 1180, y + 8, 10,
              {0: BLUE, 1: GREEN, 2: RED}[index], "", size=5)
        label(slide, f"SUMMER_T_b4_head{index}", 1200, y - 2, 70, 22, head,
              size=8.5, color=color, bold=True, align=PP_ALIGN.LEFT)
        label(slide, f"SUMMER_T_b4_tail{index}", 1268, y - 2, 250, 22, tail,
              size=8.5, align=PP_ALIGN.LEFT)

    out = HERE / "mphf_net_rebuild.pptx"
    presentation.save(out)
    _write_manifests(pyramid)
    return out


def _face_bounds(x, y, w, h, depth=8):
    """px 坐标 -> pt 期望面 bounds [x, y, w, h]。"""
    def pt_list(px_x, px_y, px_w, px_h):
        return [round(v * SCALE, 3) for v in (px_x, px_y, px_w, px_h)]
    return {
        "top": pt_list(x, y, w, depth),
        "front": pt_list(x, y + depth, w, h),
        "right": pt_list(x + w, y + depth, depth, h),
    }


def _write_manifests(pyramid):
    import json

    cuboids = []
    for index, (x, y, w, h) in enumerate(pyramid, start=1):
        cuboids.append({
            "id": f"bb{index}",
            "front": f"SUMMER_E_bb{index}_front",
            "top": f"SUMMER_E_bb{index}_top",
            "right": f"SUMMER_E_bb{index}_right",
            "expected_face_bounds_pt": _face_bounds(x, y, w, h),
            "bounds_tolerance_pt": 1.0,
        })
    for key in ("p3", "p4", "p5"):
        for prefix, tag in (("SUMMER_E_out_", "out_"),
                            ("SUMMER_E_neck_", "neck_")):
            cuboids.append({
                "id": f"{tag}{key}",
                "front": f"{prefix}{key}_front",
                "top": f"{prefix}{key}_top",
                "right": f"{prefix}{key}_right",
            })
    layering = {
        "layering_audit": {
            "cuboids": cuboids,
            "size_order": ["bb1", "bb2", "bb3", "bb4"],
            "z_order": ["bb1", "bb2", "bb3", "bb4"],
            "overlap_pairs": [["bb1", "bb2"], ["bb2", "bb3"],
                              ["bb3", "bb4"]],
        }
    }
    routes = [
        ("SUMMER_L_input_backbone", "SUMMER_E_panel_backbone", "left"),
        ("SUMMER_L_backbone_neck", "SUMMER_E_panel_neck", "left"),
        ("SUMMER_L_neck_head", "SUMMER_E_panel_head", "left"),
        ("SUMMER_L_b1_b2", "SUMMER_E_panel_b2", "left"),
        ("SUMMER_L_b2_b3", "SUMMER_E_panel_b3", "left"),
        ("SUMMER_L_neck_add1_fusion", "SUMMER_E_neck_fusion1", "left"),
        ("SUMMER_L_neck_add2_fusion", "SUMMER_E_neck_fusion2", "left"),
        ("SUMMER_L_head_iou_dec", "SUMMER_E_head_dec", "top"),
        ("SUMMER_L_head_dec_out", "SUMMER_E_head_out", "top"),
    ]
    routing = {
        "routing_audit": {
            "routes": [
                {"connector": connector, "target": target,
                 "target_edge": edge, "tolerance_pt": 1.5}
                for connector, target, edge in routes
            ],
            # 窄豁免：箭头按源图语义精确触及 UP/Down/+ 圆形徽章边缘，
            # 徽章内的单字标注文本框与线端点必然相接，非误压正文
            "ignore_text_shapes": [
                "SUMMER_E_neck_up1_text", "SUMMER_E_neck_up2_text",
                "SUMMER_E_neck_add_td1_text", "SUMMER_E_neck_add_td2_text",
                "SUMMER_E_neck_add_p5_text", "SUMMER_E_neck_add_p4_text",
                "SUMMER_E_neck_add_p3_text", "SUMMER_E_neck_down1_text",
                "SUMMER_E_neck_down2_text",
            ],
        }
    }
    (HERE / "layering_manifest.json").write_text(
        json.dumps(layering, ensure_ascii=False, indent=2), encoding="utf-8")
    (HERE / "routing_manifest.json").write_text(
        json.dumps(routing, ensure_ascii=False, indent=2), encoding="utf-8")


if __name__ == "__main__":
    print(build())

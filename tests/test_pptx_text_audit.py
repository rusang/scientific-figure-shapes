from __future__ import annotations

import importlib.util
import sys
import tempfile
import unittest
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Pt


ROOT = Path(__file__).resolve().parents[1]
MODULE_PATH = ROOT / "scripts" / "pptx_text_audit.py"


def load_module():
    if not MODULE_PATH.is_file():
        raise AssertionError("缺少 scripts/pptx_text_audit.py")
    spec = importlib.util.spec_from_file_location("pptx_text_audit", MODULE_PATH)
    assert spec and spec.loader
    module = importlib.util.module_from_spec(spec)
    sys.modules[spec.name] = module
    spec.loader.exec_module(module)
    return module


class PptxTextAuditTests(unittest.TestCase):
    def setUp(self) -> None:
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)

    def tearDown(self) -> None:
        self.temp.cleanup()

    def _presentation(self, *, text: str, size: float, align) -> Path:
        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        box = slide.shapes.add_textbox(Pt(20), Pt(20), Pt(180), Pt(42))
        box.name = "FIG_T_title"
        box.text_frame.clear()
        box.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
        paragraph = box.text_frame.paragraphs[0]
        paragraph.alignment = align
        run = paragraph.add_run()
        run.text = text
        run.font.name = "Arial"
        run.font.size = Pt(size)
        run.font.color.rgb = RGBColor(10, 20, 30)
        path = self.root / "text.pptx"
        presentation.save(path)
        return path

    def test_matching_text_typography_and_alignment_pass(self) -> None:
        module = load_module()
        path = self._presentation(text="Backbone: CSP–MEEM", size=18, align=PP_ALIGN.CENTER)
        manifest = {
            "text_audit": {
                "items": [
                    {
                        "name": "FIG_T_title",
                        "text": "Backbone: CSP–MEEM",
                        "min_font_pt": 17,
                        "align": "center",
                        "valign": "middle",
                        "max_lines": 1,
                    }
                ]
            }
        }

        result = module.audit_presentation(str(path), manifest=manifest)

        self.assertTrue(result["passed"])
        self.assertEqual(result["issues"], [])

    def test_text_content_font_alignment_and_line_count_are_audited(self) -> None:
        module = load_module()
        path = self._presentation(text="Backbone\nCSP", size=10, align=PP_ALIGN.LEFT)
        manifest = {
            "items": [
                {
                    "name": "FIG_T_title",
                    "text": "Backbone: CSP–MEEM",
                    "min_font_pt": 17,
                    "align": "center",
                    "max_lines": 1,
                }
            ]
        }

        result = module.audit_presentation(str(path), manifest=manifest)
        codes = {issue["code"] for issue in result["issues"]}

        self.assertFalse(result["passed"])
        self.assertEqual(
            {"text_mismatch", "font_too_small", "alignment_mismatch", "too_many_lines"},
            codes,
        )

    def test_missing_required_text_shape_fails(self) -> None:
        module = load_module()
        path = self._presentation(text="A", size=12, align=PP_ALIGN.CENTER)
        result = module.audit_presentation(
            str(path), manifest={"items": [{"name": "FIG_T_missing", "required": True}]}
        )
        self.assertFalse(result["passed"])
        self.assertEqual(result["issues"][0]["code"], "text_shape_missing")


if __name__ == "__main__":
    unittest.main()

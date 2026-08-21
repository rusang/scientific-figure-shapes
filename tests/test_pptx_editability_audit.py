from __future__ import annotations

import importlib.util
import sys
import tempfile
import unittest
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Pt


ROOT = Path(__file__).resolve().parents[1]
MODULE_PATH = ROOT / "scripts" / "pptx_editability_audit.py"


def load_module():
    if not MODULE_PATH.is_file():
        raise AssertionError("缺少 scripts/pptx_editability_audit.py")
    spec = importlib.util.spec_from_file_location("pptx_editability_audit", MODULE_PATH)
    assert spec and spec.loader
    module = importlib.util.module_from_spec(spec)
    sys.modules[spec.name] = module
    spec.loader.exec_module(module)
    return module


class PptxEditabilityAuditTests(unittest.TestCase):
    def setUp(self) -> None:
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)
        self.image = self.root / "image.png"
        Image.new("RGB", (320, 180), (220, 230, 240)).save(self.image)

    def tearDown(self) -> None:
        self.temp.cleanup()

    def test_full_slide_unregistered_raster_fails(self) -> None:
        module = load_module()
        presentation = Presentation()
        presentation.slide_width = Pt(320)
        presentation.slide_height = Pt(180)
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        picture = slide.shapes.add_picture(
            str(self.image), Pt(0), Pt(0), width=Pt(320), height=Pt(180)
        )
        picture.name = "FIG_R_flattened_slide"
        path = self.root / "flat.pptx"
        presentation.save(path)

        result = module.audit_presentation(str(path))

        self.assertFalse(result["passed"])
        self.assertGreater(result["slides"][0]["unregistered_raster_ratio"], 0.99)
        self.assertEqual(result["issues"][0]["code"], "unregistered_raster_area")

    def test_small_registered_crop_with_editable_shapes_passes(self) -> None:
        module = load_module()
        presentation = Presentation()
        presentation.slide_width = Pt(320)
        presentation.slide_height = Pt(180)
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Pt(0), Pt(0), Pt(320), Pt(180))
        shape.name = "FIG_E_background"
        picture = slide.shapes.add_picture(
            str(self.image), Pt(20), Pt(20), width=Pt(20), height=Pt(20)
        )
        picture.name = "FIG_R_preserved_icon"
        path = self.root / "hybrid.pptx"
        presentation.save(path)
        manifest = {
            "editability_audit": {
                "allowed_raster_shapes": ["FIG_R_preserved_icon"],
                "max_unregistered_raster_ratio": 0.01,
                "max_preserved_raster_ratio": 0.02,
            }
        }

        result = module.audit_presentation(str(path), manifest=manifest)

        self.assertTrue(result["passed"])
        self.assertEqual(result["issues"], [])


if __name__ == "__main__":
    unittest.main()

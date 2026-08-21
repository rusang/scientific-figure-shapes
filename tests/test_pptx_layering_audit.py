from __future__ import annotations

import importlib.util
import tempfile
import unittest
from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE
from pptx.dml.color import RGBColor
from pptx.util import Pt


ROOT = Path(__file__).resolve().parents[1]
MODULE_PATH = ROOT / "scripts" / "pptx_layering_audit.py"


def load_module():
    if not MODULE_PATH.is_file():
        raise AssertionError("缺少 scripts/pptx_layering_audit.py")
    spec = importlib.util.spec_from_file_location("pptx_layering_audit", MODULE_PATH)
    module = importlib.util.module_from_spec(spec)
    assert spec and spec.loader
    spec.loader.exec_module(module)
    return module


class LayeringAuditTests(unittest.TestCase):
    def setUp(self) -> None:
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)

    def tearDown(self) -> None:
        self.temp.cleanup()

    @staticmethod
    def _face(slide, name, x, y, w, h, color):
        shape = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, Pt(x), Pt(y), Pt(w), Pt(h)
        )
        shape.name = name
        shape.fill.solid()
        shape.fill.fore_color.rgb = RGBColor(*color)
        return shape

    def _add_cuboid(self, slide, prefix, x, y, w, h, colors=None):
        colors = colors or {
            "front": (180, 210, 250),
            "top": (220, 235, 255),
            "right": (120, 165, 225),
        }
        self._face(slide, f"{prefix}_front", x, y + 10, w, h, colors["front"])
        self._face(slide, f"{prefix}_top", x, y, w, 10, colors["top"])
        self._face(slide, f"{prefix}_right", x + w, y + 10, 10, h, colors["right"])

    def _manifest(self):
        return {
            "cuboids": [
                {
                    "id": "small",
                    "front": "SUMMER_E_small_front",
                    "top": "SUMMER_E_small_top",
                    "right": "SUMMER_E_small_right",
                },
                {
                    "id": "large",
                    "front": "SUMMER_E_large_front",
                    "top": "SUMMER_E_large_top",
                    "right": "SUMMER_E_large_right",
                },
            ],
            "size_order": ["small", "large"],
            "z_order": ["large", "small"],
        }

    def test_valid_layered_cuboids_pass(self) -> None:
        module = load_module()
        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        self._add_cuboid(slide, "SUMMER_E_large", 80, 80, 80, 70)
        self._add_cuboid(slide, "SUMMER_E_small", 100, 40, 40, 35)
        path = self.root / "valid.pptx"
        presentation.save(path)
        result = module.audit_presentation(str(path), manifest=self._manifest())
        self.assertTrue(result["passed"])

    def test_missing_face_fails(self) -> None:
        module = load_module()
        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        self._add_cuboid(slide, "SUMMER_E_large", 80, 80, 80, 70)
        self._face(slide, "SUMMER_E_small_front", 100, 50, 40, 35, (180, 210, 250))
        self._face(slide, "SUMMER_E_small_top", 100, 40, 40, 10, (220, 235, 255))
        path = self.root / "missing.pptx"
        presentation.save(path)
        result = module.audit_presentation(str(path), manifest=self._manifest())
        self.assertFalse(result["passed"])
        self.assertEqual(result["missing_faces"][0]["face"], "right")

    def test_flat_same_color_faces_fail(self) -> None:
        module = load_module()
        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        flat = {face: (180, 210, 250) for face in ("front", "top", "right")}
        self._add_cuboid(slide, "SUMMER_E_large", 80, 80, 80, 70, flat)
        self._add_cuboid(slide, "SUMMER_E_small", 100, 40, 40, 35)
        path = self.root / "flat.pptx"
        presentation.save(path)
        result = module.audit_presentation(str(path), manifest=self._manifest())
        self.assertFalse(result["passed"])
        self.assertEqual(result["face_color_errors"][0]["cuboid"], "large")

    def test_two_near_identical_faces_fail_despite_contrasting_third(self) -> None:
        module = load_module()
        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        near_flat = {
            "front": (225, 238, 255),
            "top": (231, 240, 253),
            "right": (158, 185, 240),
        }
        self._add_cuboid(slide, "SUMMER_E_large", 80, 80, 80, 70, near_flat)
        self._add_cuboid(slide, "SUMMER_E_small", 100, 40, 40, 35)
        path = self.root / "near-flat.pptx"
        presentation.save(path)
        result = module.audit_presentation(str(path), manifest=self._manifest())
        self.assertFalse(result["passed"])
        self.assertEqual(result["face_color_errors"][0]["cuboid"], "large")
        self.assertEqual(
            result["face_color_errors"][0]["code"], "face_contrast_too_low"
        )

    def test_wrong_z_order_fails(self) -> None:
        module = load_module()
        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        self._add_cuboid(slide, "SUMMER_E_small", 100, 40, 40, 35)
        self._add_cuboid(slide, "SUMMER_E_large", 80, 80, 80, 70)
        path = self.root / "z-order.pptx"
        presentation.save(path)
        result = module.audit_presentation(str(path), manifest=self._manifest())
        self.assertFalse(result["passed"])
        self.assertEqual(result["z_order_errors"][0]["before"], "large")

    def test_unknown_manifest_id_fails_instead_of_silent_skip(self) -> None:
        module = load_module()
        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        self._add_cuboid(slide, "SUMMER_E_small", 100, 40, 40, 35)
        self._add_cuboid(slide, "SUMMER_E_large", 80, 80, 80, 70)
        path = self.root / "typo-id.pptx"
        presentation.save(path)
        manifest = self._manifest()
        manifest["z_order"] = ["large", "smal"]
        manifest["overlap_pairs"] = [["large", "smal"]]
        result = module.audit_presentation(str(path), manifest=manifest)
        self.assertFalse(result["passed"])
        codes = [error["code"] for error in result["z_order_errors"]]
        self.assertIn("unknown_cuboid_id", codes)
        self.assertEqual(result["overlap_errors"][0]["code"], "unknown_cuboid_id")
        self.assertEqual(result["overlap_errors"][0]["unknown"], ["smal"])

    def test_cross_slide_pair_reported_not_compared(self) -> None:
        module = load_module()
        presentation = Presentation()
        slide_one = presentation.slides.add_slide(presentation.slide_layouts[6])
        slide_two = presentation.slides.add_slide(presentation.slide_layouts[6])
        self._add_cuboid(slide_one, "SUMMER_E_small", 100, 40, 40, 35)
        self._add_cuboid(slide_two, "SUMMER_E_large", 80, 80, 80, 70)
        path = self.root / "cross-slide.pptx"
        presentation.save(path)
        manifest = self._manifest()
        manifest["overlap_pairs"] = [["large", "small"]]
        result = module.audit_presentation(str(path), manifest=manifest)
        self.assertFalse(result["passed"])
        self.assertEqual(
            [error["code"] for error in result["z_order_errors"]],
            ["cross_slide_pair"],
        )
        self.assertEqual(
            [error["code"] for error in result["size_order_errors"]],
            ["cross_slide_pair"],
        )
        self.assertEqual(result["overlap_errors"][0]["code"], "cross_slide_pair")

    def test_overlap_uses_face_rects_not_union_bbox_phantom_corner(self) -> None:
        module = load_module()
        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        # cuboid a：union bbox (100,100)-(190,180)，右上角 (180,100)-(190,110)
        # 为无面幽灵区
        self._face(slide, "SUMMER_E_a_front", 100, 110, 80, 70, (180, 210, 250))
        self._face(slide, "SUMMER_E_a_top", 100, 100, 80, 10, (220, 235, 255))
        self._face(slide, "SUMMER_E_a_right", 180, 110, 10, 70, (120, 165, 225))
        # cuboid b：front 底边仅侵入 a 的幽灵区，与 a 任何面零交集，
        # 但两 union bbox 相交 8x2=16 pt^2
        self._face(slide, "SUMMER_E_b_front", 182, 60, 40, 42, (180, 210, 250))
        self._face(slide, "SUMMER_E_b_top", 182, 50, 40, 10, (220, 235, 255))
        self._face(slide, "SUMMER_E_b_right", 222, 60, 10, 42, (120, 165, 225))
        path = self.root / "phantom-overlap.pptx"
        presentation.save(path)
        manifest = {
            "cuboids": [
                {
                    "id": "a",
                    "front": "SUMMER_E_a_front",
                    "top": "SUMMER_E_a_top",
                    "right": "SUMMER_E_a_right",
                },
                {
                    "id": "b",
                    "front": "SUMMER_E_b_front",
                    "top": "SUMMER_E_b_top",
                    "right": "SUMMER_E_b_right",
                },
            ],
            "overlap_pairs": [["a", "b"]],
        }
        result = module.audit_presentation(str(path), manifest=manifest)
        self.assertFalse(result["passed"])
        self.assertEqual(
            result["overlap_errors"][0]["code"], "required_overlap_missing"
        )

    def test_real_face_overlap_still_satisfies_pair(self) -> None:
        module = load_module()
        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        self._add_cuboid(slide, "SUMMER_E_large", 80, 80, 80, 70)
        self._add_cuboid(slide, "SUMMER_E_small", 100, 40, 40, 35)
        path = self.root / "real-overlap.pptx"
        presentation.save(path)
        manifest = self._manifest()
        manifest["overlap_pairs"] = [["large", "small"]]
        result = module.audit_presentation(str(path), manifest=manifest)
        self.assertTrue(result["passed"])
        self.assertEqual(result["overlap_errors"], [])

    def test_expected_face_bounds_mismatch_fails(self) -> None:
        module = load_module()
        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        self._add_cuboid(slide, "SUMMER_E_large", 80, 80, 80, 70)
        self._add_cuboid(slide, "SUMMER_E_small", 100, 40, 40, 35)
        path = self.root / "geometry-reference.pptx"
        presentation.save(path)
        manifest = self._manifest()
        manifest["cuboids"][0]["expected_face_bounds_pt"] = {
            "front": [110, 50, 40, 35],
            "top": [100, 40, 40, 10],
            "right": [140, 50, 10, 35],
        }
        manifest["cuboids"][0]["bounds_tolerance_pt"] = 0.5
        result = module.audit_presentation(str(path), manifest=manifest)
        self.assertFalse(result["passed"])
        self.assertEqual(result["geometry_reference_errors"][0]["face"], "front")

    def test_required_gradient_face_fails_when_solid(self) -> None:
        module = load_module()
        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        self._add_cuboid(slide, "SUMMER_E_large", 80, 80, 80, 70)
        self._add_cuboid(slide, "SUMMER_E_small", 100, 40, 40, 35)
        path = self.root / "gradient-required.pptx"
        presentation.save(path)
        manifest = self._manifest()
        manifest["cuboids"][0]["gradient_faces"] = ["front"]
        result = module.audit_presentation(str(path), manifest=manifest)
        self.assertFalse(result["passed"])
        self.assertEqual(result["gradient_errors"][0]["face"], "front")

    def test_expected_face_vertices_detect_wrong_depth_geometry(self) -> None:
        module = load_module()
        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        self._face(slide, "FIG_E_cube_front", 100, 80, 60, 50, (180, 210, 250))
        self._face(slide, "FIG_E_cube_top", 100, 70, 60, 10, (220, 235, 255))
        self._face(slide, "FIG_E_cube_right", 160, 80, 10, 50, (120, 165, 225))
        path = self.root / "wrong-depth.pptx"
        presentation.save(path)
        manifest = {
            "cuboids": [
                {
                    "id": "cube",
                    "front": "FIG_E_cube_front",
                    "top": "FIG_E_cube_top",
                    "right": "FIG_E_cube_right",
                    "expected_face_vertices_pt": {
                        "top": [[100, 70], [110, 60], [170, 60], [160, 70]]
                    },
                    "vertex_tolerance_pt": 0.5,
                }
            ]
        }

        result = module.audit_presentation(str(path), manifest=manifest)

        self.assertFalse(result["passed"])
        self.assertEqual(
            result["geometry_reference_errors"][0]["code"],
            "face_vertices_mismatch",
        )

    def test_expected_gradient_colors_and_angle_are_audited(self) -> None:
        module = load_module()
        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        front = self._face(slide, "FIG_E_cube_front", 100, 80, 60, 50, (180, 210, 250))
        front.fill.gradient()
        front.fill.gradient_stops[0].color.rgb = RGBColor(220, 235, 255)
        front.fill.gradient_stops[1].color.rgb = RGBColor(120, 165, 225)
        front.fill.gradient_angle = 90
        self._face(slide, "FIG_E_cube_top", 100, 70, 60, 10, (230, 242, 255))
        self._face(slide, "FIG_E_cube_right", 160, 80, 10, 50, (90, 135, 205))
        path = self.root / "gradient-reference.pptx"
        presentation.save(path)
        manifest = {
            "cuboids": [
                {
                    "id": "cube",
                    "front": "FIG_E_cube_front",
                    "top": "FIG_E_cube_top",
                    "right": "FIG_E_cube_right",
                    "expected_gradients": {
                        "front": {
                            "colors": [[220, 235, 255], [20, 30, 40]],
                            "angle": 45,
                            "color_tolerance": 3,
                            "angle_tolerance": 2,
                        }
                    },
                }
            ]
        }

        result = module.audit_presentation(str(path), manifest=manifest)
        codes = {error["code"] for error in result["gradient_errors"]}

        self.assertFalse(result["passed"])
        self.assertEqual(codes, {"gradient_color_mismatch", "gradient_angle_mismatch"})

    def test_custom_freeform_vertices_match_reference_geometry(self) -> None:
        module = load_module()
        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        self._face(slide, "FIG_E_cube_front", 100, 70, 60, 50, (180, 210, 250))
        builder = slide.shapes.build_freeform(100, 70, scale=Pt(1))
        builder.add_line_segments([(110, 60), (170, 60), (160, 70)], close=True)
        top = builder.convert_to_shape()
        top.name = "FIG_E_cube_top"
        top.fill.solid()
        top.fill.fore_color.rgb = RGBColor(225, 240, 255)
        self._face(slide, "FIG_E_cube_right", 160, 70, 10, 50, (120, 165, 225))
        path = self.root / "custom-vertices.pptx"
        presentation.save(path)
        manifest = {
            "cuboids": [
                {
                    "id": "cube",
                    "front": "FIG_E_cube_front",
                    "top": "FIG_E_cube_top",
                    "right": "FIG_E_cube_right",
                    "expected_face_vertices_pt": {
                        "top": [[100, 70], [110, 60], [170, 60], [160, 70]]
                    },
                    "vertex_tolerance_pt": 0.5,
                }
            ]
        }

        result = module.audit_presentation(str(path), manifest=manifest)

        self.assertTrue(result["passed"], result["geometry_reference_errors"])


if __name__ == "__main__":
    unittest.main()

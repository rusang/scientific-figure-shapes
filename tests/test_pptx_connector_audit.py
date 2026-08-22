from __future__ import annotations

import importlib.util
import tempfile
import unittest
from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.util import Pt


ROOT = Path(__file__).resolve().parents[1]
MODULE_PATH = ROOT / "scripts" / "pptx_connector_audit.py"


def load_module():
    if not MODULE_PATH.is_file():
        raise AssertionError("缺少 scripts/pptx_connector_audit.py")
    spec = importlib.util.spec_from_file_location("pptx_connector_audit", MODULE_PATH)
    module = importlib.util.module_from_spec(spec)
    assert spec and spec.loader
    spec.loader.exec_module(module)
    return module


class ConnectorAuditTests(unittest.TestCase):
    def setUp(self) -> None:
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)

    def tearDown(self) -> None:
        self.temp.cleanup()

    def _save_fixture(
        self,
        name: str,
        line_points: tuple[float, float, float, float],
        text_value: str = "融合",
    ) -> Path:
        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])

        target = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE, Pt(200), Pt(100), Pt(100), Pt(40)
        )
        target.name = "SUMMER_E_fusion_box"

        text = slide.shapes.add_textbox(Pt(200), Pt(100), Pt(100), Pt(40))
        text.name = "SUMMER_E_fusion_text"
        text.text_frame.text = text_value

        x1, y1, x2, y2 = line_points
        connector = slide.shapes.add_connector(
            1, Pt(x1), Pt(y1), Pt(x2), Pt(y2)
        )
        connector.name = "SUMMER_L_branch_to_fusion"

        path = self.root / name
        presentation.save(path)
        return path

    def test_flags_connector_crossing_text_box(self) -> None:
        module = load_module()
        path = self._save_fixture("crossing.pptx", (100, 120, 350, 120))
        result = module.audit_presentation(str(path), shrink_pt=1.0)
        self.assertFalse(result["passed"])
        self.assertEqual(result["collisions"][0]["connector"], "SUMMER_L_branch_to_fusion")
        self.assertEqual(result["collisions"][0]["text_shape"], "SUMMER_E_fusion_text")

    def test_allows_connector_ending_at_text_box_boundary(self) -> None:
        module = load_module()
        path = self._save_fixture("boundary.pptx", (100, 120, 200, 120))
        result = module.audit_presentation(str(path), shrink_pt=1.0)
        self.assertTrue(result["passed"])
        self.assertEqual(result["collisions"], [])

    def test_manifest_checks_required_target_edge(self) -> None:
        module = load_module()
        path = self._save_fixture("top-edge.pptx", (250, 50, 250, 100))
        top_manifest = {
            "routes": [{
                "connector": "SUMMER_L_branch_to_fusion",
                "target": "SUMMER_E_fusion_box",
                "target_edge": "top",
                "tolerance_pt": 1.0,
            }]
        }
        result = module.audit_presentation(str(path), manifest=top_manifest)
        self.assertTrue(result["passed"])

        left_manifest = {
            "routes": [{
                "connector": "SUMMER_L_branch_to_fusion",
                "target": "SUMMER_E_fusion_box",
                "target_edge": "left",
                "tolerance_pt": 1.0,
            }]
        }
        result = module.audit_presentation(str(path), manifest=left_manifest)
        self.assertFalse(result["passed"])
        self.assertEqual(result["route_errors"][0]["code"], "target_edge_mismatch")

    def test_ignores_decorative_ellipsis_text(self) -> None:
        module = load_module()
        path = self._save_fixture(
            "ellipsis.pptx", (100, 120, 350, 120), text_value="••••"
        )
        result = module.audit_presentation(str(path), shrink_pt=1.0)
        self.assertTrue(result["passed"])
        self.assertEqual(result["collisions"], [])

    def test_elbow_connector_reported_unsupported_not_misjudged(self) -> None:
        module = load_module()
        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        text = slide.shapes.add_textbox(Pt(180), Pt(110), Pt(60), Pt(30))
        text.name = "SUMMER_E_label_text"
        text.text_frame.text = "标签"
        connector = slide.shapes.add_connector(
            MSO_CONNECTOR.ELBOW, Pt(100), Pt(100), Pt(300), Pt(160)
        )
        connector.name = "SUMMER_L_elbow_route"
        path = self.root / "elbow.pptx"
        presentation.save(path)

        result = module.audit_presentation(str(path), shrink_pt=1.0)
        self.assertFalse(result["passed"])
        self.assertEqual(result["collisions"], [])
        self.assertEqual(
            result["unsupported_connectors"][0]["connector"],
            "SUMMER_L_elbow_route",
        )
        self.assertNotIn(
            result["unsupported_connectors"][0]["preset"], ("line", "straightConnector1")
        )

    def test_flags_duplicate_connector_segments(self) -> None:
        module = load_module()
        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        for index in range(2):
            connector = slide.shapes.add_connector(
                1, Pt(100), Pt(100), Pt(200), Pt(100)
            )
            connector.name = f"SUMMER_L_duplicate_{index + 1}"
        path = self.root / "duplicate.pptx"
        presentation.save(path)

        result = module.audit_presentation(str(path))
        self.assertFalse(result["passed"])
        self.assertEqual(len(result["duplicate_segments"]), 1)

    def test_flags_degenerate_zero_length_connector(self) -> None:
        module = load_module()
        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        connector = slide.shapes.add_connector(
            1, Pt(100), Pt(100), Pt(100), Pt(100)
        )
        connector.name = "SUMMER_L_zero_length"
        path = self.root / "zero-length.pptx"
        presentation.save(path)

        result = module.audit_presentation(str(path))
        self.assertFalse(result["passed"])
        self.assertEqual(
            result["degenerate_segments"][0]["connector"],
            "SUMMER_L_zero_length",
        )

    def test_default_detection_accepts_project_specific_l_prefix(self) -> None:
        module = load_module()
        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        target = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE, Pt(200), Pt(100), Pt(100), Pt(40)
        )
        target.name = "FIG_E_target"
        connector = slide.shapes.add_connector(
            MSO_CONNECTOR.STRAIGHT, Pt(100), Pt(120), Pt(200), Pt(120)
        )
        connector.name = "FIG_L_to_target"
        path = self.root / "project-prefix.pptx"
        presentation.save(path)
        manifest = {
            "routes": [
                {
                    "connector": "FIG_L_to_target",
                    "target": "FIG_E_target",
                    "target_edge": "left",
                }
            ]
        }

        result = module.audit_presentation(str(path), manifest=manifest)

        self.assertTrue(result["passed"])
        self.assertEqual(result["connector_count"], 1)

    def test_segment_manifest_checks_dash_and_orientation(self) -> None:
        module = load_module()
        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        connector = slide.shapes.add_connector(
            MSO_CONNECTOR.STRAIGHT, Pt(100), Pt(120), Pt(200), Pt(120)
        )
        connector.name = "FIG_L_cross_scale_segment"
        path = self.root / "segment-style.pptx"
        presentation.save(path)
        manifest = {
            "segments": [
                {
                    "connector": "FIG_L_cross_scale_segment",
                    "dash": True,
                    "orientation": "vertical",
                    "tolerance_pt": 0.5,
                }
            ]
        }

        result = module.audit_presentation(str(path), manifest=manifest)
        codes = {error["code"] for error in result["style_errors"]}

        self.assertFalse(result["passed"])
        self.assertEqual(codes, {"dash_style_mismatch", "orientation_mismatch"})


if __name__ == "__main__":
    unittest.main()

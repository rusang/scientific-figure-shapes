from __future__ import annotations

import importlib.util
import sys
import tempfile
import unittest
from pathlib import Path

from PIL import Image, ImageDraw
from pptx import Presentation


ROOT = Path(__file__).resolve().parents[1]
MODULE_PATH = ROOT / "scripts" / "vector_trace.py"


def load_module():
    if not MODULE_PATH.is_file():
        raise AssertionError("缺少 scripts/vector_trace.py")
    spec = importlib.util.spec_from_file_location("vector_trace", MODULE_PATH)
    assert spec and spec.loader
    module = importlib.util.module_from_spec(spec)
    sys.modules[spec.name] = module
    spec.loader.exec_module(module)
    return module


class VectorTraceTests(unittest.TestCase):
    def setUp(self) -> None:
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)
        self.image = self.root / "icon.png"
        source = Image.new("RGB", (80, 60), "white")
        ImageDraw.Draw(source).polygon(
            [(15, 45), (15, 25), (35, 25), (35, 15), (60, 35), (35, 55), (35, 45)],
            fill=(0, 50, 180),
        )
        source.save(self.image)

    def tearDown(self) -> None:
        self.temp.cleanup()

    def test_trace_emits_editable_paths_and_svg(self) -> None:
        module = load_module()
        svg = self.root / "icon.svg"
        result = module.trace_image(
            self.image,
            threshold=245,
            min_area=20,
            epsilon_ratio=0.01,
            svg_path=svg,
        )

        self.assertTrue(result["passed"])
        self.assertGreaterEqual(len(result["paths"]), 1)
        self.assertGreaterEqual(len(result["paths"][0]["points_px"]), 6)
        self.assertTrue(svg.is_file())
        self.assertIn("<path", svg.read_text(encoding="utf-8"))

    def test_trace_materializes_freeforms_in_pptx(self) -> None:
        module = load_module()
        result = module.trace_image(self.image, threshold=245, min_area=20)
        output = self.root / "icon.pptx"
        module.materialize_pptx(result, output, width_pt=160, height_pt=120)

        presentation = Presentation(output)
        names = [shape.name for shape in presentation.slides[0].shapes]
        self.assertTrue(any(name.startswith("TRACE_E_path_") for name in names))

    def test_transparent_ring_preserves_hole_without_tracing_canvas(self) -> None:
        module = load_module()
        ring_path = self.root / "ring.png"
        ring = Image.new("RGBA", (80, 80), (0, 0, 0, 0))
        draw = ImageDraw.Draw(ring)
        draw.ellipse((10, 10, 70, 70), fill=(0, 70, 180, 255))
        draw.ellipse((28, 28, 52, 52), fill=(0, 0, 0, 0))
        ring.save(ring_path)

        result = module.trace_image(ring_path, threshold=245, min_area=20)

        self.assertTrue(result["passed"])
        self.assertLess(result["paths"][0]["area_px2"], 5000)
        self.assertTrue(any(path["hole"] for path in result["paths"]))


if __name__ == "__main__":
    unittest.main()

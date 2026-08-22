from __future__ import annotations

import importlib.util
import sys
import tempfile
import unittest
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.enum.dml import MSO_FILL_TYPE


ROOT = Path(__file__).resolve().parents[1]
MODULE_PATH = ROOT / "scripts" / "office_shape_canvas.py"


def load_module():
    if not MODULE_PATH.is_file():
        raise AssertionError("缺少 scripts/office_shape_canvas.py")
    spec = importlib.util.spec_from_file_location("office_shape_canvas", MODULE_PATH)
    assert spec and spec.loader
    module = importlib.util.module_from_spec(spec)
    sys.modules[spec.name] = module
    spec.loader.exec_module(module)
    return module


class OfficeShapeCanvasTests(unittest.TestCase):
    def setUp(self) -> None:
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)

    def tearDown(self) -> None:
        self.temp.cleanup()

    def test_pptx_and_vba_materializations_preserve_gradient_and_names(self) -> None:
        module = load_module()
        canvas = module.ShapeCanvas(320, 180)
        block = canvas.rect(
            20,
            20,
            80,
            60,
            fill=(85, 135, 220),
            line=(30, 90, 190),
            name="FIG_E_backbone_block",
        )
        canvas.gradient_fill(block, (185, 215, 255), (75, 125, 215), angle=90)
        canvas.text(
            120,
            30,
            100,
            30,
            "C5 20×C",
            name="FIG_T_c5_label",
        )
        pptx_path = self.root / "figure.pptx"
        vba_path = self.root / "figure.bas"
        canvas.save(pptx_path)
        canvas.emit_vba(vba_path)

        presentation = Presentation(pptx_path)
        shapes = {shape.name: shape for shape in presentation.slides[0].shapes}
        self.assertEqual(shapes["FIG_E_backbone_block"].fill.type, MSO_FILL_TYPE.GRADIENT)
        self.assertIn("FIG_T_c5_label", shapes)
        vba = vba_path.read_text(encoding="utf-8")
        self.assertIn("Public Sub BuildFinal()", vba)
        self.assertIn("TwoColorGradient", vba)

    def test_reference_slide_is_preserved_as_a_separate_slide(self) -> None:
        module = load_module()
        image_path = self.root / "reference.png"
        Image.new("RGB", (160, 90), (245, 245, 245)).save(image_path)
        canvas = module.ShapeCanvas(320, 180)
        canvas.add_reference_slide(image_path)
        output = self.root / "with-reference.pptx"
        canvas.save(output)

        presentation = Presentation(output)
        self.assertEqual(len(presentation.slides), 2)
        self.assertEqual(
            presentation.slides[1].shapes[0].name,
            "FIG_R_reference_full",
        )

    def test_logical_group_and_incremental_update_are_materialized(self) -> None:
        module = load_module()
        canvas = module.ShapeCanvas(320, 180)
        with canvas.group("backbone"):
            canvas.rect(
                20,
                20,
                80,
                60,
                fill=(180, 210, 250),
                line=(30, 90, 190),
                name="FIG_E_backbone",
            )
            canvas.text(20, 85, 80, 20, "C5", name="FIG_T_backbone")
        canvas.update_shape(
            "FIG_E_backbone", x=30, fill=(120, 165, 225), rotation=5
        )
        output = self.root / "updated.pptx"
        vba = self.root / "updated.bas"
        canvas.save(output)
        canvas.emit_vba(vba)

        presentation = Presentation(output)
        shapes = {shape.name: shape for shape in presentation.slides[0].shapes}
        self.assertAlmostEqual(shapes["FIG_E_backbone"].left.pt, 30, places=2)
        self.assertAlmostEqual(shapes["FIG_E_backbone"].rotation, 5, places=2)
        self.assertEqual(
            canvas.group_members("backbone"),
            ["FIG_E_backbone", "FIG_T_backbone"],
        )
        scene = canvas.scene_manifest()
        self.assertEqual(scene["groups"]["backbone"][0], "FIG_E_backbone")
        self.assertIn("Shapes.Range(Array", vba.read_text(encoding="utf-8"))

    def test_freeform_and_line_updates_keep_vba_scene_in_sync(self) -> None:
        module = load_module()
        canvas = module.ShapeCanvas(200, 120)
        canvas.freeform(
            [(10, 10), (30, 10), (30, 30), (10, 30)],
            fill=(100, 150, 220),
            name="FIG_E_polygon",
        )
        canvas.line(10, 50, 60, 50, name="FIG_L_route", arrow_end=True)

        canvas.update_shape("FIG_E_polygon", x=20, y=15, w=40, h=30)
        canvas.update_shape("FIG_L_route", x1=20, y1=60, x2=90, y2=60)

        polygon = next(op for op in canvas.ops if op.values["name"] == "FIG_E_polygon")
        route = next(op for op in canvas.ops if op.values["name"] == "FIG_L_route")
        xs = [point[0] for point in polygon.values["points"]]
        ys = [point[1] for point in polygon.values["points"]]
        self.assertEqual((min(xs), min(ys), max(xs), max(ys)), (20, 15, 60, 45))
        self.assertEqual(
            [route.values[key] for key in ("x1", "y1", "x2", "y2")],
            [20, 60, 90, 60],
        )
        vba = self.root / "incremental.bas"
        canvas.emit_vba(vba)
        payload = vba.read_text(encoding="utf-8")
        self.assertIn("20.000, 60.000, 90.000, 60.000", payload)

    def test_generated_shapes_disable_theme_shadow_effects(self) -> None:
        module = load_module()
        canvas = module.ShapeCanvas(200, 120)
        shape = canvas.round_rect(
            20,
            20,
            80,
            50,
            fill=(252, 252, 254),
            line=(190, 200, 220),
            name="FIG_E_panel",
        )

        style = shape._element.find(
            "{http://schemas.openxmlformats.org/presentationml/2006/main}style"
        )
        self.assertIsNotNone(style)
        effect_ref = style.find(
            "{http://schemas.openxmlformats.org/drawingml/2006/main}effectRef"
        )
        self.assertIsNotNone(effect_ref)
        self.assertEqual(effect_ref.get("idx"), "0")
        self.assertFalse(shape.shadow.inherit)


if __name__ == "__main__":
    unittest.main()

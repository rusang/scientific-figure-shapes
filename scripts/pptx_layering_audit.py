#!/usr/bin/env python3
"""Audit multi-face cuboid depth, size hierarchy, overlap, and z-order in PPTX."""
from __future__ import annotations

import argparse
import json
from pathlib import Path
from typing import Any


EMU_PER_POINT = 12700
FACE_NAMES = ("front", "top", "right")


def _load_pptx():
    try:
        from pptx import Presentation
    except ImportError as exc:
        raise RuntimeError(
            "pptx_layering_audit.py 需要 python-pptx：pip install python-pptx"
        ) from exc
    return Presentation


def _manifest_data(manifest: str | dict[str, Any]) -> dict[str, Any]:
    if isinstance(manifest, dict):
        data = manifest
    else:
        data = json.loads(Path(manifest).read_text(encoding="utf-8"))
    if not isinstance(data, dict):
        raise ValueError("layering manifest 顶层必须是 JSON object")
    return data.get("layering_audit", data)


def _bounds(shape) -> tuple[float, float, float, float]:
    left = shape.left / EMU_PER_POINT
    top = shape.top / EMU_PER_POINT
    right = (shape.left + shape.width) / EMU_PER_POINT
    bottom = (shape.top + shape.height) / EMU_PER_POINT
    return min(left, right), min(top, bottom), max(left, right), max(top, bottom)


def _union(rects: list[tuple[float, float, float, float]]
           ) -> tuple[float, float, float, float]:
    return (
        min(rect[0] for rect in rects),
        min(rect[1] for rect in rects),
        max(rect[2] for rect in rects),
        max(rect[3] for rect in rects),
    )


def _area(rect: tuple[float, float, float, float]) -> float:
    return max(0.0, rect[2] - rect[0]) * max(0.0, rect[3] - rect[1])


def _intersection_area(
    a: tuple[float, float, float, float],
    b: tuple[float, float, float, float],
) -> float:
    return max(0.0, min(a[2], b[2]) - max(a[0], b[0])) * max(
        0.0, min(a[3], b[3]) - max(a[1], b[1])
    )


def _rect_gap(
    a: tuple[float, float, float, float],
    b: tuple[float, float, float, float],
) -> float:
    gap_x = max(0.0, max(a[0], b[0]) - min(a[2], b[2]))
    gap_y = max(0.0, max(a[1], b[1]) - min(a[3], b[3]))
    return (gap_x * gap_x + gap_y * gap_y) ** 0.5


def _parse_rgb(value) -> tuple[int, int, int] | None:
    if value is None:
        return None
    text = str(value)
    if len(text) != 6:
        return None
    try:
        return tuple(int(text[index:index + 2], 16) for index in (0, 2, 4))  # type: ignore[return-value]
    except ValueError:
        return None


def _fill_colors(shape) -> list[tuple[int, int, int]]:
    colors: list[tuple[int, int, int]] = []
    try:
        for stop in shape.fill.gradient_stops:
            parsed = _parse_rgb(stop.color.rgb)
            if parsed is not None:
                colors.append(parsed)
    except Exception:
        pass
    if colors:
        return colors
    try:
        parsed = _parse_rgb(shape.fill.fore_color.rgb)
    except Exception:
        parsed = None
    return [parsed] if parsed is not None else []


def _has_gradient(shape) -> bool:
    namespace = "{http://schemas.openxmlformats.org/drawingml/2006/main}gradFill"
    return shape._element.spPr.find(namespace) is not None


def _min_pairwise_color_delta(colors: list[tuple[int, int, int]]) -> int:
    """每对颜色取最大通道差，再取所有对的最小值——最难分辨的一对说了算。"""
    minimum: int | None = None
    for index, first in enumerate(colors):
        for second in colors[index + 1:]:
            delta = max(abs(a - b) for a, b in zip(first, second))
            minimum = delta if minimum is None else min(minimum, delta)
    return 0 if minimum is None else minimum


def audit_presentation(
    pptx_path: str,
    *,
    manifest: str | dict[str, Any],
) -> dict[str, Any]:
    source = Path(pptx_path)
    if not source.is_file():
        raise FileNotFoundError(source)
    config = _manifest_data(manifest)
    cuboid_specs = config.get("cuboids", [])
    if not isinstance(cuboid_specs, list) or not cuboid_specs:
        raise ValueError("layering manifest 必须包含非空 cuboids 列表")

    min_face_color_delta = int(config.get("min_face_color_delta", 8))
    face_gap_tolerance = float(config.get("face_gap_tolerance_pt", 1.0))
    min_size_growth = float(config.get("min_size_growth", 1.02))
    min_overlap_area = float(config.get("min_overlap_area_pt2", 1.0))

    Presentation = _load_pptx()
    presentation = Presentation(str(source))
    indexed: dict[str, list[tuple[int, int, Any]]] = {}
    for slide_index, slide in enumerate(presentation.slides, start=1):
        for z_index, shape in enumerate(slide.shapes):
            indexed.setdefault(shape.name, []).append((slide_index, z_index, shape))

    missing_faces: list[dict[str, Any]] = []
    face_color_errors: list[dict[str, Any]] = []
    face_geometry_errors: list[dict[str, Any]] = []
    size_order_errors: list[dict[str, Any]] = []
    z_order_errors: list[dict[str, Any]] = []
    overlap_errors: list[dict[str, Any]] = []
    geometry_reference_errors: list[dict[str, Any]] = []
    gradient_errors: list[dict[str, Any]] = []
    cuboids: dict[str, dict[str, Any]] = {}

    for spec in cuboid_specs:
        cuboid_id = str(spec.get("id", "")).strip()
        faces: dict[str, tuple[int, int, Any]] = {}
        for face in FACE_NAMES:
            shape_name = str(spec.get(face, ""))
            matches = indexed.get(shape_name, [])
            if len(matches) != 1:
                missing_faces.append({
                    "cuboid": cuboid_id,
                    "face": face,
                    "shape": shape_name,
                    "matches": len(matches),
                })
            else:
                faces[face] = matches[0]
        if len(faces) != 3:
            continue
        slides = {item[0] for item in faces.values()}
        if len(slides) != 1:
            face_geometry_errors.append({
                "cuboid": cuboid_id,
                "code": "faces_cross_slide",
            })
            continue
        rects = {face: _bounds(item[2]) for face, item in faces.items()}
        palettes = {face: _fill_colors(item[2]) for face, item in faces.items()}
        colors = [palette[0] for palette in palettes.values() if palette]
        if len(colors) != 3:
            face_color_errors.append({
                "cuboid": cuboid_id,
                "code": "face_color_missing",
            })
        elif _min_pairwise_color_delta(colors) < min_face_color_delta:
            face_color_errors.append({
                "cuboid": cuboid_id,
                "code": "face_contrast_too_low",
                "colors": colors,
            })
        for face in spec.get("gradient_faces", []):
            face = str(face)
            if face not in faces:
                continue
            if not _has_gradient(faces[face][2]):
                gradient_errors.append({
                    "cuboid": cuboid_id,
                    "face": face,
                    "code": "required_gradient_missing",
                })
        expected_bounds = spec.get("expected_face_bounds_pt", {})
        bounds_tolerance = float(spec.get("bounds_tolerance_pt", 1.0))
        for face, expected in expected_bounds.items():
            if face not in rects or not isinstance(expected, list) or len(expected) != 4:
                continue
            actual_rect = rects[face]
            actual = [
                actual_rect[0], actual_rect[1],
                actual_rect[2] - actual_rect[0],
                actual_rect[3] - actual_rect[1],
            ]
            if any(abs(float(a) - float(b)) > bounds_tolerance for a, b in zip(actual, expected)):
                geometry_reference_errors.append({
                    "cuboid": cuboid_id,
                    "face": face,
                    "expected_bounds_pt": expected,
                    "actual_bounds_pt": [round(value, 3) for value in actual],
                    "tolerance_pt": bounds_tolerance,
                })
        for first, second in (("front", "top"), ("front", "right"), ("top", "right")):
            if _rect_gap(rects[first], rects[second]) > face_gap_tolerance:
                face_geometry_errors.append({
                    "cuboid": cuboid_id,
                    "code": "faces_detached",
                    "faces": [first, second],
                })
        union = _union(list(rects.values()))
        z_values = [item[1] for item in faces.values()]
        cuboids[cuboid_id] = {
            "slide": next(iter(slides)),
            "bounds_pt": union,
            "face_bounds_pt": {face: list(rect) for face, rect in rects.items()},
            "area_pt2": _area(union),
            "z_min": min(z_values),
            "z_max": max(z_values),
        }

    size_order = [str(value) for value in config.get("size_order", [])]
    for name in size_order:
        if name not in cuboids:
            size_order_errors.append({"code": "unknown_cuboid_id", "id": name})
    for smaller, larger in zip(size_order, size_order[1:]):
        if smaller not in cuboids or larger not in cuboids:
            continue
        if cuboids[smaller]["slide"] != cuboids[larger]["slide"]:
            size_order_errors.append({
                "code": "cross_slide_pair",
                "pair": [smaller, larger],
            })
            continue
        small_area = cuboids[smaller]["area_pt2"]
        large_area = cuboids[larger]["area_pt2"]
        if large_area < small_area * min_size_growth:
            size_order_errors.append({
                "smaller": smaller,
                "larger": larger,
                "small_area_pt2": round(small_area, 3),
                "large_area_pt2": round(large_area, 3),
            })

    z_order = [str(value) for value in config.get("z_order", [])]
    for name in z_order:
        if name not in cuboids:
            z_order_errors.append({"code": "unknown_cuboid_id", "id": name})
    for before, after in zip(z_order, z_order[1:]):
        if before not in cuboids or after not in cuboids:
            continue
        if cuboids[before]["slide"] != cuboids[after]["slide"]:
            z_order_errors.append({
                "code": "cross_slide_pair",
                "pair": [before, after],
            })
            continue
        if cuboids[before]["z_max"] >= cuboids[after]["z_min"]:
            z_order_errors.append({
                "before": before,
                "after": after,
                "before_z_max": cuboids[before]["z_max"],
                "after_z_min": cuboids[after]["z_min"],
            })

    for pair in config.get("overlap_pairs", []):
        if not isinstance(pair, list) or len(pair) != 2:
            overlap_errors.append({"code": "overlap_pair_invalid", "pair": pair})
            continue
        first, second = str(pair[0]), str(pair[1])
        unknown = [name for name in (first, second) if name not in cuboids]
        if unknown:
            overlap_errors.append({
                "code": "unknown_cuboid_id",
                "pair": [first, second],
                "unknown": unknown,
            })
            continue
        if cuboids[first]["slide"] != cuboids[second]["slide"]:
            overlap_errors.append({
                "code": "cross_slide_pair",
                "pair": [first, second],
            })
            continue
        # 按面矩形对求最大交集：三面 union bbox 的空角（top-right 幽灵区）
        # 不算真实视觉重叠；斜面仍以其自身 bbox 近似。
        overlap = max(
            _intersection_area(tuple(rect_a), tuple(rect_b))
            for rect_a in cuboids[first]["face_bounds_pt"].values()
            for rect_b in cuboids[second]["face_bounds_pt"].values()
        )
        if overlap < min_overlap_area:
            overlap_errors.append({
                "code": "required_overlap_missing",
                "pair": [first, second],
                "overlap_area_pt2": round(overlap, 3),
            })

    passed = not any((
        missing_faces,
        face_color_errors,
        face_geometry_errors,
        size_order_errors,
        z_order_errors,
        overlap_errors,
        geometry_reference_errors,
        gradient_errors,
    ))
    return {
        "passed": passed,
        "pptx": str(source.resolve()),
        "cuboids": cuboids,
        "missing_faces": missing_faces,
        "face_color_errors": face_color_errors,
        "face_geometry_errors": face_geometry_errors,
        "size_order_errors": size_order_errors,
        "z_order_errors": z_order_errors,
        "overlap_errors": overlap_errors,
        "geometry_reference_errors": geometry_reference_errors,
        "gradient_errors": gradient_errors,
    }


def main() -> int:
    parser = argparse.ArgumentParser(
        description="Audit PPTX cuboid faces, depth contrast, hierarchy, overlap, and z-order."
    )
    parser.add_argument("pptx")
    parser.add_argument("--manifest", required=True)
    parser.add_argument("--json-out")
    parser.add_argument("--pretty", action="store_true")
    args = parser.parse_args()
    try:
        result = audit_presentation(args.pptx, manifest=args.manifest)
    except Exception as exc:
        result = {"passed": False, "runtime_error": str(exc)}
    payload = json.dumps(result, ensure_ascii=False, indent=2 if args.pretty else None)
    print(payload)
    if args.json_out:
        target = Path(args.json_out)
        target.parent.mkdir(parents=True, exist_ok=True)
        target.write_text(payload + "\n", encoding="utf-8")
    return 0 if result["passed"] else 1


if __name__ == "__main__":
    raise SystemExit(main())

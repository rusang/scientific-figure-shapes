from __future__ import annotations

from contextlib import contextmanager
from dataclasses import dataclass
import importlib.util
import json
from pathlib import Path
from typing import Iterable, Sequence

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Pt


Color = tuple[int, int, int]


def hex_rgb(value: str) -> Color:
    value = value.lstrip("#")
    return tuple(int(value[i : i + 2], 16) for i in (0, 2, 4))  # type: ignore[return-value]


@dataclass
class RecordedOp:
    kind: str
    values: dict


class ShapeCanvas:
    """Small PowerPoint drawing layer that also records an equivalent VBA scene."""

    def __init__(
        self,
        width_pt: float,
        height_pt: float,
        font_name: str = "PingFang SC",
        name_prefix: str = "FIG",
    ) -> None:
        self.width_pt = width_pt
        self.height_pt = height_pt
        self.font_name = font_name
        self.name_prefix = name_prefix.rstrip("_")
        self.prs = Presentation()
        self.prs.slide_width = Pt(width_pt)
        self.prs.slide_height = Pt(height_pt)
        self.slide = self.prs.slides.add_slide(self.prs.slide_layouts[6])
        self.ops: list[RecordedOp] = []
        self._counter = 0
        self._active_group: str | None = None
        self._groups: dict[str, list[str]] = {}
        self._shape_index: dict[str, object] = {}

    def _name(self, stem: str) -> str:
        self._counter += 1
        return f"{self.name_prefix}_{stem}_{self._counter:03d}"

    def _record(self, op: RecordedOp, shape=None) -> None:
        name = str(op.values.get("name", ""))
        if name and name in self._shape_index:
            raise ValueError(f"duplicate shape name: {name}")
        if self._active_group:
            op.values["group"] = self._active_group
            self._groups.setdefault(self._active_group, []).append(name)
        self.ops.append(op)
        if name and shape is not None:
            self._shape_index[name] = shape

    @contextmanager
    def group(self, group_id: str):
        """Assign subsequently created shapes to a stable logical group."""
        group_id = group_id.strip()
        if not group_id:
            raise ValueError("group_id cannot be empty")
        previous = self._active_group
        self._active_group = group_id
        self._groups.setdefault(group_id, [])
        try:
            yield self
        finally:
            self._active_group = previous

    def group_members(self, group_id: str) -> list[str]:
        return list(self._groups.get(group_id, []))

    def update_shape(self, name: str, **changes) -> None:
        """Incrementally update an existing shape and its recorded VBA scene."""
        shape = self._shape_index.get(name)
        if shape is None:
            raise KeyError(name)
        op = next((item for item in self.ops if item.values.get("name") == name), None)
        if op is None:
            raise KeyError(name)
        allowed = {
            "x", "y", "w", "h", "x1", "y1", "x2", "y2",
            "rotation", "fill", "line", "weight", "transparency", "value",
        }
        unknown = set(changes) - allowed
        if unknown:
            raise ValueError(f"unsupported update fields: {sorted(unknown)}")
        endpoint_keys = {"x1", "y1", "x2", "y2"}
        if op.kind == "line":
            if set(changes) & {"x", "y", "w", "h", "fill", "value"}:
                raise ValueError("line updates use x1/y1/x2/y2 and line/weight")
            endpoints = {
                key: float(changes.get(key, op.values[key])) for key in endpoint_keys
            }
            for key, value in endpoints.items():
                op.values[key] = value
            shape.left = Pt(min(endpoints["x1"], endpoints["x2"]))
            shape.top = Pt(min(endpoints["y1"], endpoints["y2"]))
            shape.width = Pt(abs(endpoints["x2"] - endpoints["x1"]))
            shape.height = Pt(abs(endpoints["y2"] - endpoints["y1"]))
            xfrm = shape._element.spPr.xfrm
            for attribute, flipped in (
                ("flipH", endpoints["x1"] > endpoints["x2"]),
                ("flipV", endpoints["y1"] > endpoints["y2"]),
            ):
                if flipped:
                    xfrm.set(attribute, "1")
                elif attribute in xfrm.attrib:
                    del xfrm.attrib[attribute]
        elif set(changes) & endpoint_keys:
            raise ValueError("x1/y1/x2/y2 are only valid for line operations")

        coordinate_attrs = {"x": "left", "y": "top", "w": "width", "h": "height"}
        coordinate_changes = set(changes) & set(coordinate_attrs)
        if op.kind == "freeform" and coordinate_changes:
            points = op.values["points"]
            old_left = min(point[0] for point in points)
            old_top = min(point[1] for point in points)
            old_width = max(point[0] for point in points) - old_left
            old_height = max(point[1] for point in points) - old_top
            if old_width <= 0 or old_height <= 0:
                raise ValueError("cannot resize a degenerate freeform")
            new_left = float(changes.get("x", old_left))
            new_top = float(changes.get("y", old_top))
            new_width = float(changes.get("w", old_width))
            new_height = float(changes.get("h", old_height))
            op.values["points"] = [
                (
                    new_left + (point[0] - old_left) / old_width * new_width,
                    new_top + (point[1] - old_top) / old_height * new_height,
                )
                for point in points
            ]
            shape.left, shape.top = Pt(new_left), Pt(new_top)
            shape.width, shape.height = Pt(new_width), Pt(new_height)
        elif op.kind != "line":
            for key, attr in coordinate_attrs.items():
                if key in changes:
                    setattr(shape, attr, Pt(float(changes[key])))
                    op.values[key] = float(changes[key])
        if "rotation" in changes:
            shape.rotation = float(changes["rotation"])
            op.values["rotation"] = float(changes["rotation"])
        if "fill" in changes:
            transparency = int(changes.get("transparency", op.values.get("transparency", 0)))
            self._set_fill(shape, changes["fill"], transparency)
            op.values["fill"] = changes["fill"]
            op.values["transparency"] = transparency
            if "gradient" in op.values:
                op.values["gradient"] = None
        if "line" in changes or "weight" in changes:
            line = changes.get("line", op.values.get("line"))
            weight = float(changes.get("weight", op.values.get("weight", 0.8)))
            self._set_line(shape, line, weight, bool(op.values.get("dash", False)))
            op.values["line"] = line
            op.values["weight"] = weight
        if "value" in changes:
            if not getattr(shape, "has_text_frame", False):
                raise ValueError(f"shape {name!r} has no text frame")
            shape.text_frame.paragraphs[0].runs[0].text = str(changes["value"])
            op.values["value"] = str(changes["value"])

    def scene_manifest(self) -> dict:
        return {
            "schema_version": "1.0",
            "canvas": {"width_pt": self.width_pt, "height_pt": self.height_pt},
            "groups": {key: list(value) for key, value in self._groups.items()},
            "operations": [
                {"kind": op.kind, **op.values}
                for op in self.ops
            ],
        }

    def export_scene_manifest(self, path: Path) -> None:
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_text(
            json.dumps(self.scene_manifest(), ensure_ascii=False, indent=2) + "\n",
            encoding="utf-8",
        )

    @staticmethod
    def _set_fill(shape, color: Color | None, transparency: int = 0) -> None:
        if color is None:
            shape.fill.background()
            return
        shape.fill.solid()
        shape.fill.fore_color.rgb = RGBColor(*color)
        shape.fill.transparency = max(0, min(100, transparency))

    @staticmethod
    def _set_line(shape, color: Color | None, weight: float, dash: bool = False) -> None:
        if color is None:
            shape.line.fill.background()
            return
        shape.line.color.rgb = RGBColor(*color)
        shape.line.width = Pt(weight)
        if dash:
            shape.line.dash_style = MSO_LINE_DASH_STYLE.DASH

    @staticmethod
    def _disable_theme_effects(shape) -> None:
        style = shape._element.find(qn("p:style"))
        if style is not None:
            effect_ref = style.find(qn("a:effectRef"))
            if effect_ref is not None:
                effect_ref.set("idx", "0")
        try:
            shape.shadow.inherit = False
        except (AttributeError, ValueError):
            pass

    def shape(
        self,
        shape_type: MSO_SHAPE,
        vba_shape: str,
        x: float,
        y: float,
        w: float,
        h: float,
        *,
        fill: Color | None = None,
        line: Color | None = None,
        weight: float = 0.8,
        transparency: int = 0,
        rotation: float = 0,
        name: str | None = None,
    ):
        shape = self.slide.shapes.add_shape(shape_type, Pt(x), Pt(y), Pt(w), Pt(h))
        shape.name = name or self._name("E_shape")
        shape.rotation = rotation
        self._set_fill(shape, fill, transparency)
        self._set_line(shape, line, weight)
        self._disable_theme_effects(shape)
        self._record(
            RecordedOp(
                "shape",
                {
                    "name": shape.name,
                    "vba_shape": vba_shape,
                    "x": x,
                    "y": y,
                    "w": w,
                    "h": h,
                    "fill": fill,
                    "line": line,
                    "weight": weight,
                    "transparency": transparency,
                    "rotation": rotation,
                    "gradient": None,
                },
            ),
            shape,
        )
        return shape

    def rect(self, x: float, y: float, w: float, h: float, **kwargs):
        return self.shape(MSO_SHAPE.RECTANGLE, "msoShapeRectangle", x, y, w, h, **kwargs)

    def round_rect(self, x: float, y: float, w: float, h: float, **kwargs):
        shape = self.shape(MSO_SHAPE.ROUNDED_RECTANGLE, "msoShapeRoundedRectangle", x, y, w, h, **kwargs)
        shape.adjustments[0] = 0.08
        return shape

    def ellipse(self, x: float, y: float, w: float, h: float, **kwargs):
        return self.shape(MSO_SHAPE.OVAL, "msoShapeOval", x, y, w, h, **kwargs)

    def cube(self, x: float, y: float, w: float, h: float, **kwargs):
        return self.shape(MSO_SHAPE.CUBE, "msoShapeCube", x, y, w, h, **kwargs)

    def right_arrow(self, x: float, y: float, w: float, h: float, **kwargs):
        return self.shape(MSO_SHAPE.RIGHT_ARROW, "msoShapeRightArrow", x, y, w, h, **kwargs)

    def pentagon(self, x: float, y: float, w: float, h: float, **kwargs):
        return self.shape(MSO_SHAPE.PENTAGON, "msoShapePentagon", x, y, w, h, **kwargs)

    def freeform(
        self,
        points: Sequence[tuple[float, float]],
        *,
        fill: Color | None = None,
        line: Color | None = None,
        weight: float = 0.8,
        transparency: int = 0,
        rotation: float = 0,
        name: str | None = None,
    ):
        if len(points) < 3:
            raise ValueError("freeform requires at least three points")
        builder = self.slide.shapes.build_freeform(
            points[0][0], points[0][1], scale=Pt(1)
        )
        builder.add_line_segments(points[1:], close=True)
        shape = builder.convert_to_shape()
        shape.name = name or self._name("E_freeform")
        shape.rotation = rotation
        self._set_fill(shape, fill, transparency)
        self._set_line(shape, line, weight)
        self._disable_theme_effects(shape)
        self._record(
            RecordedOp(
                "freeform",
                {
                    "name": shape.name,
                    "points": list(points),
                    "fill": fill,
                    "line": line,
                    "weight": weight,
                    "transparency": transparency,
                    "rotation": rotation,
                    "gradient": None,
                },
            ),
            shape,
        )
        return shape

    def gradient_fill(
        self,
        shape,
        start_color: Color,
        end_color: Color,
        *,
        angle: float = 90,
    ) -> None:
        shape.fill.gradient()
        stops = shape.fill.gradient_stops
        stops[0].position = 0.0
        stops[0].color.rgb = RGBColor(*start_color)
        stops[1].position = 1.0
        stops[1].color.rgb = RGBColor(*end_color)
        shape.fill.gradient_angle = float(angle)
        for op in reversed(self.ops):
            if op.values.get("name") == shape.name:
                op.values["gradient"] = {
                    "start": start_color,
                    "end": end_color,
                    "angle": float(angle),
                }
                return
        raise ValueError(f"shape {shape.name!r} is not recorded")

    def text(
        self,
        x: float,
        y: float,
        w: float,
        h: float,
        value: str,
        *,
        size: float = 9,
        color: Color = (0, 0, 0),
        bold: bool = False,
        align: str = "center",
        valign: str = "middle",
        margin: float = 0,
        font_name: str | None = None,
        rotation: float = 0,
        name: str | None = None,
    ):
        box = self.slide.shapes.add_textbox(Pt(x), Pt(y), Pt(w), Pt(h))
        box.name = name or self._name("E_text")
        box.rotation = rotation
        self._disable_theme_effects(box)
        tf = box.text_frame
        tf.clear()
        tf.word_wrap = True
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = Pt(margin)
        tf.vertical_anchor = {
            "top": MSO_ANCHOR.TOP,
            "middle": MSO_ANCHOR.MIDDLE,
            "bottom": MSO_ANCHOR.BOTTOM,
        }[valign]
        p = tf.paragraphs[0]
        p.alignment = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT}[align]
        run = p.add_run()
        run.text = value
        run.font.name = font_name or self.font_name
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = RGBColor(*color)
        self._record(
            RecordedOp(
                "text",
                {
                    "name": box.name,
                    "x": x,
                    "y": y,
                    "w": w,
                    "h": h,
                    "value": value,
                    "size": size,
                    "color": color,
                    "bold": bold,
                    "align": align,
                    "valign": valign,
                    "margin": margin,
                    "font_name": font_name or self.font_name,
                    "rotation": rotation,
                },
            ),
            box,
        )
        return box

    @staticmethod
    def _append_arrowhead(connector, begin: bool, style: str) -> None:
        line = connector._element.spPr.get_or_add_ln()
        tag = "a:headEnd" if begin else "a:tailEnd"
        node = OxmlElement(tag)
        node.set("type", style)
        node.set("w", "sm")
        node.set("len", "sm")
        line.append(node)

    def line(
        self,
        x1: float,
        y1: float,
        x2: float,
        y2: float,
        *,
        color: Color = (0, 0, 0),
        weight: float = 1,
        dash: bool = False,
        arrow_end: bool = False,
        arrow_begin: bool = False,
        name: str | None = None,
    ):
        conn = self.slide.shapes.add_connector(
            MSO_CONNECTOR.STRAIGHT, Pt(x1), Pt(y1), Pt(x2), Pt(y2)
        )
        conn.name = name or self._name("L_line")
        self._set_line(conn, color, weight, dash)
        if arrow_end:
            self._append_arrowhead(conn, False, "triangle")
        if arrow_begin:
            self._append_arrowhead(conn, True, "triangle")
        self._record(
            RecordedOp(
                "line",
                {
                    "name": conn.name,
                    "x1": x1,
                    "y1": y1,
                    "x2": x2,
                    "y2": y2,
                    "color": color,
                    "weight": weight,
                    "dash": dash,
                    "arrow_end": arrow_end,
                    "arrow_begin": arrow_begin,
                },
            ),
            conn,
        )
        return conn

    def polyline(
        self,
        points: Sequence[tuple[float, float]],
        *,
        color: Color = (0, 0, 0),
        weight: float = 1,
        dash: bool = False,
        arrow_end: bool = True,
    ) -> None:
        for index, ((x1, y1), (x2, y2)) in enumerate(zip(points, points[1:])):
            self.line(
                x1,
                y1,
                x2,
                y2,
                color=color,
                weight=weight,
                dash=dash,
                arrow_end=arrow_end and index == len(points) - 2,
            )

    def picture(
        self,
        path: Path,
        x: float,
        y: float,
        w: float,
        h: float,
        *,
        name: str | None = None,
    ):
        pic = self.slide.shapes.add_picture(str(path), Pt(x), Pt(y), width=Pt(w), height=Pt(h))
        pic.name = name or self._name("R_picture")
        pic.lock_aspect_ratio = True
        self._record(
            RecordedOp(
                "picture",
                {"name": pic.name, "path": str(path.resolve()), "x": x, "y": y, "w": w, "h": h},
            ),
            pic,
        )
        return pic

    def add_reference_slide(self, path: Path, *, name: str | None = None):
        name = name or f"{self.name_prefix}_R_reference_full"
        slide = self.prs.slides.add_slide(self.prs.slide_layouts[6])
        pic = slide.shapes.add_picture(
            str(path), Pt(0), Pt(0), width=Pt(self.width_pt), height=Pt(self.height_pt)
        )
        pic.name = name
        pic.lock_aspect_ratio = True
        self._record(
            RecordedOp(
                "reference_slide_picture",
                {
                    "name": name,
                    "path": str(path.resolve()),
                    "x": 0.0,
                    "y": 0.0,
                    "w": self.width_pt,
                    "h": self.height_pt,
                },
            ),
            pic,
        )
        return slide

    def grid(
        self,
        x: float,
        y: float,
        w: float,
        h: float,
        rows: int,
        cols: int,
        *,
        fill: Color,
        line: Color,
        transparency: int = 35,
    ) -> None:
        cw, ch = w / cols, h / rows
        for row in range(rows):
            for col in range(cols):
                self.rect(
                    x + col * cw,
                    y + row * ch,
                    cw,
                    ch,
                    fill=fill,
                    line=line,
                    weight=0.45,
                    transparency=transparency,
                )

    def save(self, path: Path) -> None:
        path.parent.mkdir(parents=True, exist_ok=True)
        self.prs.save(path)

    @staticmethod
    def _vba_color(color: Color | None) -> str:
        if color is None:
            return "-1"
        return f"RGB({color[0]}, {color[1]}, {color[2]})"

    @staticmethod
    def _vba_string(value: str) -> str:
        parts = value.replace("\r", "").split("\n")
        encoded = [f'"{part.replace(chr(34), chr(34) * 2)}"' for part in parts]
        return " & vbLf & ".join(encoded)

    def emit_vba(self, path: Path) -> None:
        module_path = Path(__file__).with_name("office_shape_vba.py")
        spec = importlib.util.spec_from_file_location("office_shape_vba_runtime", module_path)
        assert spec and spec.loader
        module = importlib.util.module_from_spec(spec)
        spec.loader.exec_module(module)
        module.emit_vba(self, path)

def add_bullets(
    canvas: ShapeCanvas,
    x: float,
    y: float,
    width: float,
    items: Iterable[str],
    *,
    size: float,
    line_height: float,
    color: Color = (0, 0, 0),
) -> None:
    for index, item in enumerate(items):
        canvas.text(
            x,
            y + index * line_height,
            width,
            line_height,
            f"•  {item}",
            size=size,
            color=color,
            align="left",
            valign="middle",
        )

#!/usr/bin/env python3
"""Measure raster dependency and editable-shape coverage in a PPTX."""

from __future__ import annotations

import argparse
import json
from pathlib import Path
from typing import Any


EMU_PER_POINT = 12700


def _load_pptx():
    try:
        from pptx import Presentation
        from pptx.enum.shapes import MSO_SHAPE_TYPE
    except ImportError as exc:
        raise RuntimeError(
            "pptx_editability_audit.py 需要 python-pptx：pip install python-pptx"
        ) from exc
    return Presentation, MSO_SHAPE_TYPE


def _manifest_data(manifest: str | dict[str, Any] | None) -> dict[str, Any]:
    if manifest is None:
        return {}
    if isinstance(manifest, dict):
        data = manifest
    else:
        data = json.loads(Path(manifest).read_text(encoding="utf-8"))
    if not isinstance(data, dict):
        raise ValueError("editability manifest 顶层必须是 JSON object")
    config = data.get("editability_audit", data)
    if not isinstance(config, dict):
        raise ValueError("editability_audit 必须是 JSON object")
    return config


def _clipped_area(shape, slide_width: int, slide_height: int) -> int:
    left = max(0, min(slide_width, shape.left))
    top = max(0, min(slide_height, shape.top))
    right = max(0, min(slide_width, shape.left + shape.width))
    bottom = max(0, min(slide_height, shape.top + shape.height))
    return max(0, right - left) * max(0, bottom - top)


def audit_presentation(
    pptx_path: str,
    *,
    manifest: str | dict[str, Any] | None = None,
) -> dict[str, Any]:
    source = Path(pptx_path)
    if not source.is_file():
        raise FileNotFoundError(source)
    config = _manifest_data(manifest)
    allowed = {str(name) for name in config.get("allowed_raster_shapes", [])}
    reference_names = {
        str(name) for name in config.get("reference_slide_shapes", [])
    }
    max_unregistered = float(config.get("max_unregistered_raster_ratio", 0.05))
    max_preserved = float(config.get("max_preserved_raster_ratio", 0.25))
    min_editable_count = int(config.get("min_editable_shape_count", 1))

    Presentation, MSO_SHAPE_TYPE = _load_pptx()
    presentation = Presentation(str(source))
    slide_area = presentation.slide_width * presentation.slide_height
    issues: list[dict[str, Any]] = []
    slide_reports: list[dict[str, Any]] = []

    for slide_index, slide in enumerate(presentation.slides, start=1):
        names = {shape.name for shape in slide.shapes}
        is_reference = bool(names & reference_names) or any(
            name.endswith("_R_reference_full") for name in names
        )
        if is_reference:
            slide_reports.append(
                {"slide": slide_index, "excluded_reference_slide": True}
            )
            continue

        editable_count = 0
        editable_area = 0
        registered_area = 0
        unregistered_area = 0
        pictures: list[dict[str, Any]] = []
        for shape in slide.shapes:
            area = _clipped_area(shape, presentation.slide_width, presentation.slide_height)
            if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                registered = shape.name in allowed
                if registered:
                    registered_area += area
                else:
                    unregistered_area += area
                pictures.append(
                    {
                        "name": shape.name,
                        "registered": registered,
                        "area_ratio": round(area / slide_area, 6) if slide_area else 0,
                    }
                )
            else:
                editable_count += 1
                editable_area += area

        unregistered_ratio = min(1.0, unregistered_area / slide_area) if slide_area else 0
        registered_ratio = min(1.0, registered_area / slide_area) if slide_area else 0
        editable_ratio = min(1.0, editable_area / slide_area) if slide_area else 0
        report = {
            "slide": slide_index,
            "excluded_reference_slide": False,
            "editable_shape_count": editable_count,
            "editable_area_ratio": round(editable_ratio, 6),
            "preserved_raster_ratio": round(registered_ratio, 6),
            "unregistered_raster_ratio": round(unregistered_ratio, 6),
            "pictures": pictures,
        }
        slide_reports.append(report)
        if unregistered_ratio > max_unregistered:
            issues.append(
                {
                    "code": "unregistered_raster_area",
                    "slide": slide_index,
                    "maximum": max_unregistered,
                    "actual": round(unregistered_ratio, 6),
                }
            )
        if editable_count < min_editable_count:
            issues.append(
                {
                    "code": "insufficient_editable_shapes",
                    "slide": slide_index,
                    "minimum": min_editable_count,
                    "actual": editable_count,
                }
            )
        if registered_ratio > max_preserved:
            issues.append(
                {
                    "code": "preserved_raster_area_too_large",
                    "slide": slide_index,
                    "maximum": max_preserved,
                    "actual": round(registered_ratio, 6),
                }
            )

    return {
        "passed": not issues,
        "pptx": str(source.resolve()),
        "slides": slide_reports,
        "issues": issues,
    }


def main() -> int:
    parser = argparse.ArgumentParser(description="Audit PPTX editability and raster dependency.")
    parser.add_argument("pptx")
    parser.add_argument("--manifest")
    parser.add_argument("--json-out")
    parser.add_argument("--pretty", action="store_true")
    args = parser.parse_args()
    try:
        result = audit_presentation(args.pptx, manifest=args.manifest)
    except Exception as exc:
        result = {"passed": False, "runtime_error": str(exc)}
    payload = json.dumps(result, ensure_ascii=False, indent=2 if args.pretty else None)
    print(payload)
    if args.json_out:
        target = Path(args.json_out)
        target.parent.mkdir(parents=True, exist_ok=True)
        target.write_text(payload + "\n", encoding="utf-8")
    return 0 if result["passed"] else 1


if __name__ == "__main__":
    raise SystemExit(main())

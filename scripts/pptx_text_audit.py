#!/usr/bin/env python3
"""Audit required text, typography, alignment, and explicit wrapping in PPTX."""

from __future__ import annotations

import argparse
import json
import math
from pathlib import Path
from typing import Any


EMU_PER_POINT = 12700


def _load_pptx():
    try:
        from pptx import Presentation
    except ImportError as exc:
        raise RuntimeError(
            "pptx_text_audit.py 需要 python-pptx：pip install python-pptx"
        ) from exc
    return Presentation


def _manifest_data(manifest: str | dict[str, Any]) -> dict[str, Any]:
    if isinstance(manifest, dict):
        data = manifest
    else:
        data = json.loads(Path(manifest).read_text(encoding="utf-8"))
    if not isinstance(data, dict):
        raise ValueError("text manifest 顶层必须是 JSON object")
    config = data.get("text_audit", data)
    if not isinstance(config, dict):
        raise ValueError("text_audit 必须是 JSON object")
    return config


def _font_sizes(shape) -> list[float]:
    sizes: list[float] = []
    if not getattr(shape, "has_text_frame", False):
        return sizes
    for paragraph in shape.text_frame.paragraphs:
        for run in paragraph.runs:
            if run.font.size is not None:
                sizes.append(run.font.size.pt)
    return sizes


def _alignment(shape) -> str:
    from pptx.enum.text import PP_ALIGN

    values = {
        PP_ALIGN.LEFT: "left",
        PP_ALIGN.CENTER: "center",
        PP_ALIGN.RIGHT: "right",
        PP_ALIGN.JUSTIFY: "justify",
        PP_ALIGN.DISTRIBUTE: "distribute",
    }
    seen = {
        values.get(paragraph.alignment, "inherited")
        for paragraph in shape.text_frame.paragraphs
        if paragraph.text or paragraph.runs
    }
    return next(iter(seen)) if len(seen) == 1 else "mixed"


def _vertical_alignment(shape) -> str:
    from pptx.enum.text import MSO_ANCHOR

    return {
        MSO_ANCHOR.TOP: "top",
        MSO_ANCHOR.MIDDLE: "middle",
        MSO_ANCHOR.BOTTOM: "bottom",
    }.get(shape.text_frame.vertical_anchor, "inherited")


def _explicit_lines(shape) -> int:
    text = shape.text_frame.text
    return max(1, text.count("\n") + 1)


def _estimated_lines(shape, font_pt: float) -> int:
    """Conservative geometry estimate; opt-in because PowerPoint wrapping varies."""
    width_pt = shape.width / EMU_PER_POINT
    if width_pt <= 0 or font_pt <= 0:
        return 1
    total = 0
    for line in shape.text_frame.text.splitlines() or [""]:
        units = sum(1.0 if ord(char) > 127 else 0.56 for char in line)
        total += max(1, math.ceil(units * font_pt / width_pt))
    return total


def audit_presentation(
    pptx_path: str,
    *,
    manifest: str | dict[str, Any],
) -> dict[str, Any]:
    source = Path(pptx_path)
    if not source.is_file():
        raise FileNotFoundError(source)
    config = _manifest_data(manifest)
    items = config.get("items", [])
    if not isinstance(items, list) or not items:
        raise ValueError("text manifest 必须包含非空 items 列表")

    Presentation = _load_pptx()
    presentation = Presentation(str(source))
    indexed: dict[str, list[tuple[int, Any]]] = {}
    for slide_index, slide in enumerate(presentation.slides, start=1):
        for shape in slide.shapes:
            if getattr(shape, "has_text_frame", False):
                indexed.setdefault(shape.name, []).append((slide_index, shape))

    issues: list[dict[str, Any]] = []
    reports: list[dict[str, Any]] = []
    for item in items:
        if not isinstance(item, dict):
            issues.append({"code": "text_item_invalid"})
            continue
        name = str(item.get("name", "")).strip()
        if not name:
            issues.append({"code": "text_name_missing"})
            continue
        matches = indexed.get(name, [])
        required = bool(item.get("required", True))
        if not matches:
            if required:
                issues.append({"code": "text_shape_missing", "name": name})
            continue
        if len(matches) != 1:
            issues.append(
                {"code": "text_shape_duplicate", "name": name, "matches": len(matches)}
            )
            continue
        slide_index, shape = matches[0]
        actual_text = shape.text_frame.text
        sizes = _font_sizes(shape)
        min_size = min(sizes) if sizes else None
        alignment = _alignment(shape)
        valign = _vertical_alignment(shape)
        lines = _explicit_lines(shape)
        report = {
            "name": name,
            "slide": slide_index,
            "text": actual_text,
            "min_font_pt": round(min_size, 3) if min_size is not None else None,
            "align": alignment,
            "valign": valign,
            "explicit_lines": lines,
        }

        expected = item.get("text")
        if expected is not None and actual_text != str(expected):
            issues.append(
                {"code": "text_mismatch", "name": name, "expected": expected, "actual": actual_text}
            )
        contains = item.get("contains")
        if contains is not None and str(contains) not in actual_text:
            issues.append(
                {"code": "text_fragment_missing", "name": name, "expected": contains}
            )
        if "min_font_pt" in item and (
            min_size is None or min_size < float(item["min_font_pt"])
        ):
            issues.append(
                {
                    "code": "font_too_small",
                    "name": name,
                    "minimum": float(item["min_font_pt"]),
                    "actual": min_size,
                }
            )
        if "max_font_pt" in item and sizes and max(sizes) > float(item["max_font_pt"]):
            issues.append(
                {
                    "code": "font_too_large",
                    "name": name,
                    "maximum": float(item["max_font_pt"]),
                    "actual": max(sizes),
                }
            )
        if "align" in item and alignment != str(item["align"]).lower():
            issues.append(
                {"code": "alignment_mismatch", "name": name, "actual": alignment}
            )
        if "valign" in item and valign != str(item["valign"]).lower():
            issues.append(
                {"code": "vertical_alignment_mismatch", "name": name, "actual": valign}
            )
        if "max_lines" in item and lines > int(item["max_lines"]):
            issues.append(
                {"code": "too_many_lines", "name": name, "maximum": int(item["max_lines"]), "actual": lines}
            )
        if bool(item.get("estimate_overflow", False)) and sizes:
            estimated = _estimated_lines(shape, max(sizes))
            capacity = max(1, int((shape.height / EMU_PER_POINT) / (max(sizes) * 1.2)))
            report["estimated_lines"] = estimated
            report["estimated_capacity"] = capacity
            if estimated > capacity:
                issues.append(
                    {"code": "estimated_text_overflow", "name": name, "estimated_lines": estimated, "capacity": capacity}
                )
        reports.append(report)

    return {
        "passed": not issues,
        "pptx": str(source.resolve()),
        "items": reports,
        "issues": issues,
    }


def main() -> int:
    parser = argparse.ArgumentParser(description="Audit PPTX text content and typography.")
    parser.add_argument("pptx")
    parser.add_argument("--manifest", required=True)
    parser.add_argument("--json-out")
    parser.add_argument("--pretty", action="store_true")
    args = parser.parse_args()
    try:
        result = audit_presentation(args.pptx, manifest=args.manifest)
    except Exception as exc:
        result = {"passed": False, "runtime_error": str(exc)}
    payload = json.dumps(result, ensure_ascii=False, indent=2 if args.pretty else None)
    print(payload)
    if args.json_out:
        target = Path(args.json_out)
        target.parent.mkdir(parents=True, exist_ok=True)
        target.write_text(payload + "\n", encoding="utf-8")
    return 0 if result["passed"] else 1


if __name__ == "__main__":
    raise SystemExit(main())

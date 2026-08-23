#!/usr/bin/env python3
"""Audit editable PPTX connector routing against text boxes and target edges."""
from __future__ import annotations

import argparse
import json
import math
import re
from pathlib import Path
from typing import Any


EMU_PER_POINT = 12700
VALID_EDGES = {"left", "right", "top", "bottom"}
DECORATIVE_TEXT = re.compile(r"^[\s•·.…⋮⋯]+$")
PRST_GEOM = "{http://schemas.openxmlformats.org/drawingml/2006/main}prstGeom"
STRAIGHT_PRSTS = {"line", "straightConnector1"}


def _load_pptx():
    try:
        from pptx import Presentation
        from pptx.enum.shapes import MSO_SHAPE_TYPE
    except ImportError as exc:
        raise RuntimeError(
            "pptx_connector_audit.py 需要 python-pptx：pip install python-pptx"
        ) from exc
    return Presentation, MSO_SHAPE_TYPE


def _points(value: int) -> float:
    return float(value) / EMU_PER_POINT


def _bounds(shape) -> tuple[float, float, float, float]:
    left = _points(shape.left)
    top = _points(shape.top)
    right = _points(shape.left + shape.width)
    bottom = _points(shape.top + shape.height)
    return min(left, right), min(top, bottom), max(left, right), max(top, bottom)


def _truthy_xml(value: str | None) -> bool:
    return value in {"1", "true", "True"}


def _connector_preset(shape) -> str:
    prst_geom = shape._element.spPr.find(PRST_GEOM)
    return prst_geom.get("prst", "") if prst_geom is not None else ""


def _is_dashed(shape) -> bool:
    namespace = "{http://schemas.openxmlformats.org/drawingml/2006/main}prstDash"
    node = shape._element.spPr.get_or_add_ln().find(namespace)
    return node is not None and node.get("val", "solid") != "solid"


def _line_endpoints(shape) -> tuple[tuple[float, float], tuple[float, float]]:
    left, top, right, bottom = _bounds(shape)
    xfrm = shape._element.spPr.xfrm
    flip_h = _truthy_xml(xfrm.get("flipH"))
    flip_v = _truthy_xml(xfrm.get("flipV"))
    x1, x2 = (right, left) if flip_h else (left, right)
    y1, y2 = (bottom, top) if flip_v else (top, bottom)
    return (x1, y1), (x2, y2)


def _shrink_rect(rect: tuple[float, float, float, float], amount: float
                 ) -> tuple[float, float, float, float] | None:
    left, top, right, bottom = rect
    shrunk = left + amount, top + amount, right - amount, bottom - amount
    if shrunk[2] <= shrunk[0] or shrunk[3] <= shrunk[1]:
        return None
    return shrunk


def _segment_intersects_rect(
    start: tuple[float, float],
    end: tuple[float, float],
    rect: tuple[float, float, float, float],
) -> bool:
    """Liang-Barsky line clipping against an axis-aligned rectangle."""
    x1, y1 = start
    x2, y2 = end
    left, top, right, bottom = rect
    dx, dy = x2 - x1, y2 - y1
    p = (-dx, dx, -dy, dy)
    q = (x1 - left, right - x1, y1 - top, bottom - y1)
    lower, upper = 0.0, 1.0
    for pi, qi in zip(p, q):
        if abs(pi) < 1e-12:
            if qi < 0:
                return False
            continue
        ratio = qi / pi
        if pi < 0:
            lower = max(lower, ratio)
        else:
            upper = min(upper, ratio)
        if lower > upper:
            return False
    return True


def _manifest_data(manifest: str | dict[str, Any] | None) -> dict[str, Any]:
    if manifest is None:
        return {}
    if isinstance(manifest, dict):
        data = manifest
    else:
        data = json.loads(Path(manifest).read_text(encoding="utf-8"))
    if not isinstance(data, dict):
        raise ValueError("routing manifest 顶层必须是 JSON object")
    return data.get("routing_audit", data)


def _edge_contains_point(
    point: tuple[float, float],
    rect: tuple[float, float, float, float],
    edge: str,
    tolerance: float,
) -> bool:
    x, y = point
    left, top, right, bottom = rect
    if edge == "left":
        return abs(x - left) <= tolerance and top - tolerance <= y <= bottom + tolerance
    if edge == "right":
        return abs(x - right) <= tolerance and top - tolerance <= y <= bottom + tolerance
    if edge == "top":
        return abs(y - top) <= tolerance and left - tolerance <= x <= right + tolerance
    if edge == "bottom":
        return abs(y - bottom) <= tolerance and left - tolerance <= x <= right + tolerance
    return False


def _point_in_rect(
    point: tuple[float, float],
    rect: tuple[float, float, float, float],
    tolerance: float,
) -> bool:
    x, y = point
    left, top, right, bottom = rect
    return (left - tolerance <= x <= right + tolerance
            and top - tolerance <= y <= bottom + tolerance)


def _point_segment_distance(
    point: tuple[float, float],
    start: tuple[float, float],
    end: tuple[float, float],
) -> float:
    px, py = point
    x1, y1 = start
    x2, y2 = end
    dx, dy = x2 - x1, y2 - y1
    length_sq = dx * dx + dy * dy
    if length_sq < 1e-12:
        return math.hypot(px - x1, py - y1)
    t = max(0.0, min(1.0, ((px - x1) * dx + (py - y1) * dy) / length_sq))
    return math.hypot(px - (x1 + t * dx), py - (y1 + t * dy))


def _endpoint_anchored(
    point: tuple[float, float],
    other: tuple[float, float],
    anchor_rects: list[tuple[float, float, float, float]],
    peer_segments: list[tuple[tuple[float, float], tuple[float, float]]],
    tolerance: float,
) -> bool:
    """端点锚定：贴/扎进某 shape，或接续另一连接线（含 trunk 中段 T 形接续）。

    整条线深居某 shape 内部时该 shape 视为背景容器，不提供锚定。
    """
    for rect in anchor_rects:
        if not _point_in_rect(point, rect, tolerance):
            continue
        inner = _shrink_rect(rect, tolerance)
        if (inner is not None and _point_in_rect(point, inner, 0.0)
                and _point_in_rect(other, inner, 0.0)):
            continue
        return True
    for seg_start, seg_end in peer_segments:
        if _point_segment_distance(point, seg_start, seg_end) <= tolerance:
            return True
    return False


def _segment_key(
    start: tuple[float, float],
    end: tuple[float, float],
    tolerance: float,
) -> tuple[tuple[int, int], tuple[int, int]]:
    scale = 1.0 / tolerance
    points = [
        (round(start[0] * scale), round(start[1] * scale)),
        (round(end[0] * scale), round(end[1] * scale)),
    ]
    points.sort()
    return points[0], points[1]


def audit_presentation(
    pptx_path: str,
    *,
    shrink_pt: float = 1.0,
    manifest: str | dict[str, Any] | None = None,
    connector_prefix: str | None = None,
    segment_tolerance_pt: float = 0.01,
) -> dict[str, Any]:
    if shrink_pt < 0:
        raise ValueError("shrink_pt must be >= 0")
    if segment_tolerance_pt <= 0:
        raise ValueError("segment_tolerance_pt must be > 0")
    source = Path(pptx_path)
    if not source.is_file():
        raise FileNotFoundError(source)

    Presentation, MSO_SHAPE_TYPE = _load_pptx()
    presentation = Presentation(str(source))
    config = _manifest_data(manifest)
    configured_prefix = config.get("connector_prefix", connector_prefix)
    if configured_prefix is not None:
        configured_prefix = str(configured_prefix)
    route_connector_names = {
        str(route.get("connector", ""))
        for route in config.get("routes", [])
        if isinstance(route, dict)
    }
    ignore_connectors = set(config.get("ignore_connectors", []))
    ignore_text_shapes = set(config.get("ignore_text_shapes", []))
    ignore_dangling = set(config.get("ignore_dangling", []))
    anchor_tolerance = float(config.get("anchor_tolerance_pt", 1.5))
    min_connector_count = config.get("min_connector_count")
    ignore_pairs = {
        tuple(pair) for pair in config.get("ignore_pairs", [])
        if isinstance(pair, list) and len(pair) == 2
    }

    collisions: list[dict[str, Any]] = []
    route_errors: list[dict[str, Any]] = []
    duplicate_segments: list[dict[str, Any]] = []
    degenerate_segments: list[dict[str, Any]] = []
    unsupported_connectors: list[dict[str, Any]] = []
    style_errors: list[dict[str, Any]] = []
    dangling_endpoints: list[dict[str, Any]] = []
    indexed_shapes: dict[str, list[tuple[int, Any]]] = {}
    connector_count = 0
    text_count = 0

    for slide_index, slide in enumerate(presentation.slides, start=1):
        connectors = []
        texts = []
        anchor_rects: list[tuple[float, float, float, float]] = []
        for shape in slide.shapes:
            indexed_shapes.setdefault(shape.name, []).append((slide_index, shape))
            if shape.shape_type != MSO_SHAPE_TYPE.LINE:
                try:
                    anchor_rects.append(_bounds(shape))
                except TypeError:
                    pass  # 继承几何（无 xfrm）的 shape 不作锚定候选
            if shape.shape_type == MSO_SHAPE_TYPE.LINE:
                named_connector = (
                    shape.name in route_connector_names
                    or (
                        shape.name.startswith(configured_prefix)
                        if configured_prefix is not None
                        else "_L_" in shape.name
                    )
                )
                if named_connector and shape.name not in ignore_connectors:
                    preset = _connector_preset(shape)
                    if preset in STRAIGHT_PRSTS:
                        connectors.append(shape)
                    else:
                        # elbow/curved：端点≠bbox 对角线，按直线审计必错，显式拒审
                        unsupported_connectors.append({
                            "slide": slide_index,
                            "connector": shape.name,
                            "preset": preset,
                        })
            elif getattr(shape, "has_text_frame", False):
                text = shape.text.strip()
                if (text and not DECORATIVE_TEXT.fullmatch(text)
                        and shape.name not in ignore_text_shapes):
                    texts.append((shape, text))

        connector_count += len(connectors)
        text_count += len(texts)
        endpoint_map = {
            connector.name: _line_endpoints(connector)
            for connector in connectors
        }
        seen_segments: dict[tuple[tuple[int, int], tuple[int, int]], tuple[str, Any]] = {}
        for connector in connectors:
            start, end = endpoint_map[connector.name]
            length = math.hypot(end[0] - start[0], end[1] - start[1])
            if length <= segment_tolerance_pt:
                degenerate_segments.append({
                    "slide": slide_index,
                    "connector": connector.name,
                    "segment_pt": [
                        round(start[0], 3), round(start[1], 3),
                        round(end[0], 3), round(end[1], 3),
                    ],
                })
            else:
                key = _segment_key(start, end, segment_tolerance_pt)
                if key in seen_segments:
                    first_name, _ = seen_segments[key]
                    duplicate_segments.append({
                        "slide": slide_index,
                        "connectors": [first_name, connector.name],
                        "segment_pt": [
                            round(start[0], 3), round(start[1], 3),
                            round(end[0], 3), round(end[1], 3),
                        ],
                    })
                else:
                    seen_segments[key] = (connector.name, connector)
                if connector.name not in ignore_dangling:
                    peers = [
                        points for name, points in endpoint_map.items()
                        if name != connector.name
                    ]
                    for tag, point, other in (("start", start, end),
                                              ("end", end, start)):
                        if not _endpoint_anchored(point, other, anchor_rects,
                                                  peers, anchor_tolerance):
                            dangling_endpoints.append({
                                "slide": slide_index,
                                "connector": connector.name,
                                "endpoint": tag,
                                "point_pt": [round(point[0], 3),
                                             round(point[1], 3)],
                            })
            for text_shape, text in texts:
                if (connector.name, text_shape.name) in ignore_pairs:
                    continue
                inner = _shrink_rect(_bounds(text_shape), shrink_pt)
                if inner is None or not _segment_intersects_rect(start, end, inner):
                    continue
                collisions.append({
                    "slide": slide_index,
                    "connector": connector.name,
                    "text_shape": text_shape.name,
                    "text": text.replace("\n", " ")[:80],
                    "segment_pt": [
                        round(start[0], 3), round(start[1], 3),
                        round(end[0], 3), round(end[1], 3),
                    ],
                })

    for route in config.get("routes", []):
        if not isinstance(route, dict):
            route_errors.append({"code": "route_invalid", "route": route})
            continue
        connector_name = str(route.get("connector", ""))
        target_name = str(route.get("target", ""))
        edge = str(route.get("target_edge", "")).lower()
        tolerance = float(route.get("tolerance_pt", 1.0))
        if edge not in VALID_EDGES:
            route_errors.append({
                "code": "target_edge_invalid",
                "connector": connector_name,
                "target_edge": edge,
            })
            continue
        connector_matches = indexed_shapes.get(connector_name, [])
        target_matches = indexed_shapes.get(target_name, [])
        if len(connector_matches) != 1 or len(target_matches) != 1:
            route_errors.append({
                "code": "route_shape_missing_or_ambiguous",
                "connector": connector_name,
                "connector_matches": len(connector_matches),
                "target": target_name,
                "target_matches": len(target_matches),
            })
            continue
        connector_slide, connector = connector_matches[0]
        target_slide, target = target_matches[0]
        preset = _connector_preset(connector)
        if preset not in STRAIGHT_PRSTS:
            route_errors.append({
                "code": "unsupported_connector_geometry",
                "connector": connector_name,
                "preset": preset,
            })
            continue
        if connector_slide != target_slide:
            route_errors.append({
                "code": "route_cross_slide",
                "connector": connector_name,
                "target": target_name,
            })
            continue
        endpoints = _line_endpoints(connector)
        target_rect = _bounds(target)
        if not any(_edge_contains_point(point, target_rect, edge, tolerance) for point in endpoints):
            route_errors.append({
                "code": "target_edge_mismatch",
                "connector": connector_name,
                "target": target_name,
                "expected_edge": edge,
                "endpoints_pt": [list(point) for point in endpoints],
                "target_bounds_pt": list(target_rect),
            })
        source_name = route.get("source")
        if source_name is not None:
            source_name = str(source_name)
            source_edge = str(route.get("source_edge", "")).lower()
            if source_edge not in VALID_EDGES:
                route_errors.append({
                    "code": "source_edge_invalid",
                    "connector": connector_name,
                    "source_edge": source_edge,
                })
                continue
            source_matches = indexed_shapes.get(source_name, [])
            if len(source_matches) != 1:
                route_errors.append({
                    "code": "source_shape_missing_or_ambiguous",
                    "connector": connector_name,
                    "source": source_name,
                    "source_matches": len(source_matches),
                })
                continue
            source_slide, source_shape = source_matches[0]
            if source_slide != connector_slide:
                route_errors.append({
                    "code": "route_cross_slide",
                    "connector": connector_name,
                    "source": source_name,
                })
                continue
            source_rect = _bounds(source_shape)
            if not any(_edge_contains_point(point, source_rect, source_edge,
                                            tolerance)
                       for point in endpoints):
                route_errors.append({
                    "code": "source_edge_mismatch",
                    "connector": connector_name,
                    "source": source_name,
                    "expected_edge": source_edge,
                    "endpoints_pt": [list(point) for point in endpoints],
                    "source_bounds_pt": list(source_rect),
                })

    for segment in config.get("segments", []):
        if not isinstance(segment, dict):
            style_errors.append({"code": "segment_invalid", "segment": segment})
            continue
        connector_name = str(segment.get("connector", ""))
        matches = indexed_shapes.get(connector_name, [])
        if len(matches) != 1:
            style_errors.append({
                "code": "segment_shape_missing_or_ambiguous",
                "connector": connector_name,
                "matches": len(matches),
            })
            continue
        _, connector = matches[0]
        preset = _connector_preset(connector)
        if preset not in STRAIGHT_PRSTS:
            style_errors.append({
                "code": "unsupported_connector_geometry",
                "connector": connector_name,
                "preset": preset,
            })
            continue
        if "dash" in segment:
            expected_dash = bool(segment["dash"])
            actual_dash = _is_dashed(connector)
            if actual_dash != expected_dash:
                style_errors.append({
                    "code": "dash_style_mismatch",
                    "connector": connector_name,
                    "expected": expected_dash,
                    "actual": actual_dash,
                })
        orientation = segment.get("orientation")
        if orientation is not None:
            orientation = str(orientation).lower()
            tolerance = float(segment.get("tolerance_pt", 0.5))
            start, end = _line_endpoints(connector)
            delta_x = abs(end[0] - start[0])
            delta_y = abs(end[1] - start[1])
            actual_orientation = (
                "horizontal" if delta_y <= tolerance
                else "vertical" if delta_x <= tolerance
                else "diagonal"
            )
            if orientation not in {"horizontal", "vertical", "diagonal"}:
                style_errors.append({
                    "code": "orientation_invalid",
                    "connector": connector_name,
                    "expected": orientation,
                })
            elif actual_orientation != orientation:
                style_errors.append({
                    "code": "orientation_mismatch",
                    "connector": connector_name,
                    "expected": orientation,
                    "actual": actual_orientation,
                })

    coverage_errors: list[dict[str, Any]] = []
    if min_connector_count is not None:
        expected = int(min_connector_count)
        if connector_count < expected:
            coverage_errors.append({
                "code": "connector_inventory_shortfall",
                "expected_min": expected,
                "actual": connector_count,
            })

    return {
        "passed": (
            not collisions and not route_errors
            and not duplicate_segments and not degenerate_segments
            and not unsupported_connectors and not style_errors
            and not dangling_endpoints and not coverage_errors
        ),
        "pptx": str(source.resolve()),
        "slides": len(presentation.slides),
        "connector_count": connector_count,
        "text_shape_count": text_count,
        "shrink_pt": shrink_pt,
        "segment_tolerance_pt": segment_tolerance_pt,
        "collisions": collisions,
        "route_errors": route_errors,
        "duplicate_segments": duplicate_segments,
        "degenerate_segments": degenerate_segments,
        "unsupported_connectors": unsupported_connectors,
        "style_errors": style_errors,
        "dangling_endpoints": dangling_endpoints,
        "coverage_errors": coverage_errors,
    }


def main() -> int:
    parser = argparse.ArgumentParser(
        description="Audit editable PPTX connectors for text collisions and target-edge routing."
    )
    parser.add_argument("pptx")
    parser.add_argument("--manifest")
    parser.add_argument("--shrink-pt", type=float, default=1.0)
    parser.add_argument(
        "--connector-prefix",
        help="optional exact project prefix; default detects stable names containing _L_",
    )
    parser.add_argument("--segment-tolerance-pt", type=float, default=0.01)
    parser.add_argument("--json-out")
    parser.add_argument("--pretty", action="store_true")
    args = parser.parse_args()

    try:
        result = audit_presentation(
            args.pptx,
            shrink_pt=args.shrink_pt,
            manifest=args.manifest,
            connector_prefix=args.connector_prefix,
            segment_tolerance_pt=args.segment_tolerance_pt,
        )
    except Exception as exc:
        result = {
            "passed": False,
            "collisions": [],
            "route_errors": [{"code": "runtime_error", "message": str(exc)}],
            "duplicate_segments": [],
            "degenerate_segments": [],
            "style_errors": [],
            "dangling_endpoints": [],
            "coverage_errors": [],
        }
    payload = json.dumps(result, ensure_ascii=False, indent=2 if args.pretty else None)
    print(payload)
    if args.json_out:
        target = Path(args.json_out)
        target.parent.mkdir(parents=True, exist_ok=True)
        target.write_text(payload + "\n", encoding="utf-8")
    return 0 if result["passed"] else 1


if __name__ == "__main__":
    raise SystemExit(main())
